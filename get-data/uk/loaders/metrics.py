"""Nạp chỉ số vào core.metric (registry) + core.metric_value: IoD 2025, Census 2021, UK HPI, Defra PCM."""
from __future__ import annotations

import csv
import io
import json
import logging
import re
import zipfile
from datetime import date, datetime

from .base import STAGE, LoadContext, LoadError, Loader

log = logging.getLogger(__name__)

MV_COLS = ["level", "code", "metric_id", "period", "value_num", "value_json"]
MV_COLS_CI = MV_COLS + ["ci_low", "ci_high"]


def register_metric(ctx: LoadContext, mid: str, *, category: str, dataset: str, value_type: str, period_type: str,
                    name: str, unit: str | None = None, decimals: int = 0, direction: int = 0,
                    native_level: str | None = None, aggregation: str = "none", scope: str | None = None,
                    sensitivity: str = "normal", scorable: bool = True, description: str | None = None,
                    sort: int = 0) -> None:
    ctx.conn.execute("""
        INSERT INTO core.metric (id, market_code, category_code, dataset_id, value_type, unit, decimals, direction,
            period_type, native_level, aggregation, comparable_scope_id, sensitivity, is_scorable, sort_order)
        VALUES (%s, 'uk', %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        ON CONFLICT (id) DO UPDATE SET category_code = EXCLUDED.category_code, dataset_id = EXCLUDED.dataset_id,
            value_type = EXCLUDED.value_type, unit = EXCLUDED.unit, decimals = EXCLUDED.decimals,
            direction = EXCLUDED.direction, period_type = EXCLUDED.period_type, native_level = EXCLUDED.native_level,
            aggregation = EXCLUDED.aggregation, comparable_scope_id = EXCLUDED.comparable_scope_id,
            sensitivity = EXCLUDED.sensitivity, is_scorable = EXCLUDED.is_scorable, sort_order = EXCLUDED.sort_order""",
        (mid, category, dataset, value_type, unit, decimals, direction, period_type, native_level, aggregation,
         scope, sensitivity, scorable and sensitivity == "normal", sort))
    ctx.conn.execute("""
        INSERT INTO core.metric_translation (metric_id, locale_code, name, description) VALUES (%s, 'en', %s, %s)
        ON CONFLICT (metric_id, locale_code) DO UPDATE SET name = EXCLUDED.name, description = EXCLUDED.description""",
        (mid, name, description))


def insert_values(ctx: LoadContext, stage: str, label: str) -> tuple[int, int]:
    """stage(level, code, metric_id, period, value_num, value_json) → core.metric_value (upsert).

    Trả về (số dòng ghi, số dòng không khớp mã khu). Mã không khớp được log theo cấp để dễ kiểm tra.
    """
    total = ctx.scalar(f"SELECT count(*) FROM {stage}")
    unmatched = ctx.conn.execute(f"""
        SELECT s.level, count(*), (array_agg(DISTINCT s.code))[1:5] FROM {stage} s
        WHERE NOT EXISTS (SELECT 1 FROM core.geo_area a WHERE a.market_code='uk' AND a.level_code=s.level AND a.code=s.code)
        GROUP BY 1""").fetchall()
    n_unmatched = sum(r[1] for r in unmatched)
    if unmatched:
        log.warning("    %s: %s giá trị không khớp mã khu (bỏ qua): %s", label, n_unmatched, unmatched)
    n = ctx.sql(f"upsert metric_value {label}", f"""
        INSERT INTO core.metric_value (market_code, area_id, metric_id, period, value_num, value_json, ci_low, ci_high,
                                       release_id)
        SELECT 'uk', a.id, s.metric_id, s.period, s.value_num, s.value_json, s.ci_low, s.ci_high, %s
        FROM {stage} s JOIN core.geo_area a ON a.market_code='uk' AND a.level_code=s.level AND a.code=s.code
        ON CONFLICT (market_code, area_id, metric_id, period) DO UPDATE SET value_num = EXCLUDED.value_num,
            value_json = EXCLUDED.value_json, ci_low = EXCLUDED.ci_low, ci_high = EXCLUDED.ci_high,
            release_id = EXCLUDED.release_id, computed_at = now()""",
        (ctx.release_id,))
    ctx.expect(f"metric_value {label} (+ không khớp)", n + n_unmatched, total)
    return n, n_unmatched


def mv_stage(ctx: LoadContext, name: str) -> str:
    return ctx.stage_table(name, MV_COLS_CI, {"period": "date", "value_num": "float8", "value_json": "jsonb",
                                               "ci_low": "float8", "ci_high": "float8"})


def num(s):
    if isinstance(s, (int, float)):
        return float(s)
    s = (s or "").strip().replace(",", "")
    if s in ("", "MISSING", "..", "x", ":", "-", "[x]", "[z]", "[c]", "[u]"):
        return None
    try:
        return float(s)
    except ValueError:
        raise LoadError(f"Giá trị số không đọc được: {s!r}")


# ======================================================================== IoD =====
IOD_DOMAINS = [  # (tiền tố cột, id, tên)
    ("Index of Multiple Deprivation (IMD)", "imd", "Index of Multiple Deprivation"),
    ("Income", "income", "Income deprivation"),
    ("Employment", "employment", "Employment deprivation"),
    ("Education, Skills and Training", "education", "Education, skills & training deprivation"),
    ("Health Deprivation and Disability", "health", "Health deprivation & disability"),
    ("Crime", "crime", "Crime (IoD domain)"),
    ("Barriers to Housing and Services", "barriers", "Barriers to housing & services"),
    ("Living Environment", "living_env", "Living environment deprivation"),
    ("Income Deprivation Affecting Children Index (IDACI)", "idaci", "Income deprivation affecting children"),
    ("Income Deprivation Affecting Older People (IDAOPI)", "idaopi", "Income deprivation affecting older people"),
    # sub-domain (cùng File 7)
    ("Children and Young People Sub-domain", "sub_children", "Children and young people education sub-domain"),
    ("Adult Skills Sub-domain", "sub_adult_skills", "Adult skills sub-domain"),
    ("Geographical Barriers Sub-domain", "sub_geo_barriers", "Geographical barriers sub-domain"),
    ("Wider Barriers Sub-domain", "sub_wider_barriers", "Wider barriers sub-domain"),
    ("Indoors Sub-domain", "sub_indoors", "Indoors living environment sub-domain"),
    ("Outdoors Sub-domain", "sub_outdoors", "Outdoors living environment sub-domain"),
]


class IodLoader(Loader):
    source = "iod2025"
    depends_on = ("boundaries_lsoa21",)
    description = "IoD 2025 File 7 → metric uk.iod25.<domain>_{score,rank,decile} theo LSOA (England)"

    def load(self, ctx: LoadContext) -> dict:
        f = ctx.file(r"File_7_.*\.csv$")
        with open(f, encoding="utf-8-sig", newline="") as fh:
            rd = csv.DictReader(fh)
            cols = rd.fieldnames
            find = {}
            for prefix, key, _ in IOD_DOMAINS:
                for kind in ("Score", "Rank", "Decile"):
                    hits = [c for c in cols if re.match(re.escape(prefix) + r"\s+" + kind, c)]
                    if len(hits) != 1:
                        raise LoadError(f"IoD: không xác định cột {prefix} {kind}: {hits}")
                    find[(key, kind)] = hits[0]
            period = date(2025, 1, 1)
            rows = []
            for r in rd:
                code = r["LSOA code (2021)"]
                for (key, kind), col in find.items():
                    rows.append(("lsoa", code, f"uk.iod25.{key}_{kind.lower()}", period, num(r[col]), None))
        for prefix, key, name in IOD_DOMAINS:
            common = dict(category="economy", dataset="uk.mhclg.iod2025", period_type="snapshot", native_level="lsoa",
                          scope="uk.england")
            register_metric(ctx, f"uk.iod25.{key}_score", value_type="index", decimals=3, direction=-1,
                            name=f"{name} score", **common)
            register_metric(ctx, f"uk.iod25.{key}_rank", value_type="rank", direction=1,
                            name=f"{name} rank (1 = most deprived)", scorable=False, **common)
            register_metric(ctx, f"uk.iod25.{key}_decile", value_type="rank", direction=1,
                            name=f"{name} decile (1 = most deprived 10%)", scorable=False, **common)
        t = mv_stage(ctx, "mv_iod")
        ctx.copy_rows(t, MV_COLS, rows, "COPY IoD values")
        n, unmatched = insert_values(ctx, t, "IoD 2025")
        ctx.drop_stage("mv_iod")
        return {"values": n, "unmatched": unmatched}


# ===================================================================== Census =====
CENSUS_LEVELS = {"oa": "oa", "lsoa": "lsoa", "msoa": "msoa", "ltla": "lad", "rgn": "region", "ctry": "nation"}
CENSUS_SENSITIVE = {"ts021", "ts030", "ts004", "ts005", "ts015", "ts019", "ts025", "ts026", "ts027", "ts032"}  # dân tộc, tôn giáo, nơi sinh: chỉ mô tả, không tính điểm (PLAN §7)


class CensusLoader(Loader):
    source = "census2021"
    depends_on = ("boundaries_oa21", "onspd")
    description = ("Census 2021 bảng TS → metric uk.census21.<ts> (value_num = tổng, value_json = phân bố theo nhóm) "
                   "ở OA/LSOA/MSOA/LAD/region/nation")

    def load(self, ctx: LoadContext) -> dict:
        period = date(2021, 3, 21)
        stats = {}
        for z in ctx.files(r"^census2021-ts.*\.zip$"):
            ts = re.search(r"(ts\d+[a-z]?)", z.name).group(1)
            t = mv_stage(ctx, "mv_census")
            title = None
            with zipfile.ZipFile(z) as zf:
                for member in zf.namelist():
                    m = re.match(rf"census2021-{ts}-(\w+)\.csv$", member)
                    if not m or m.group(1) not in CENSUS_LEVELS:
                        continue
                    level = CENSUS_LEVELS[m.group(1)]
                    with zf.open(member) as f:
                        rd = csv.reader(io.TextIOWrapper(f, encoding="utf-8-sig", newline=""))
                        header = next(rd, None)
                        if not header:  # Nomis đôi khi để file rỗng ở một cấp (ví dụ ts010 LSOA)
                            log.warning("    %s rỗng, bỏ qua cấp %s", member, level)
                            continue
                        cats = header[3:]
                        if not cats:
                            raise LoadError(f"{member}: không có cột số liệu")
                        title = title or cats[0].split(":")[0].strip()
                        labels = [c.split(";")[0].strip() for c in cats]

                        def rows():
                            for r in rd:
                                vals = [num(v) for v in r[3:]]
                                # cột đầu thường là Total; phần còn lại là phân bố (giữ nguyên nhãn đầy đủ)
                                dist = {lab.split(":", 1)[-1].strip(): v for lab, v in zip(labels[1:], vals[1:])}
                                yield (level, r[2], f"uk.census21.{ts}", period, vals[0], json.dumps(dist))
                        ctx.copy_rows(t, MV_COLS, rows(), f"COPY {member}")
            sensitive = ts in CENSUS_SENSITIVE
            register_metric(ctx, f"uk.census21.{ts}", category="people", dataset="uk.ons.census2021",
                            value_type="distribution", period_type="snapshot", native_level="oa", aggregation="sum",
                            scope="uk.ew", name=f"Census 2021: {title}",
                            sensitivity="sensitive" if sensitive else "normal", scorable=not sensitive,
                            description=f"Nomis bulk {ts.upper()}; value_num = tổng, value_json = số lượng theo nhóm")
            n, unmatched = insert_values(ctx, t, f"Census {ts}")
            stats[ts] = {"values": n, "unmatched": unmatched}
            ctx.drop_stage("mv_census")
        if not stats:
            raise LoadError("Không có file census2021-ts*.zip")
        stats["lad_filled_from_msoa"] = fill_census_lad_from_msoa(ctx)
        return stats


CENSUS_LAD_FILL_SQL = """
    -- LAD hiện hành chưa có giá trị Census (LAD mới sau 2021, ví dụ Cumberland, North Yorkshire, Somerset):
    -- cộng dồn từ MSOA con (liên kết msoa → lad từ ONSPD). Số liệu Census là số đếm nên cộng được.
    WITH lad_missing AS (
        SELECT l.id AS lad_id, m.id AS metric_id
        FROM core.geo_area l CROSS JOIN core.metric m
        WHERE l.market_code = 'uk' AND l.level_code = 'lad' AND l.is_active AND left(l.code, 1) IN ('E', 'W')
          AND m.id LIKE 'uk.census21.%%' AND m.id <> 'uk.census21.ts006'  -- mật độ không cộng được, xử lý riêng
          AND NOT EXISTS (SELECT 1 FROM core.metric_value v WHERE v.market_code = 'uk' AND v.area_id = l.id
                          AND v.metric_id = m.id)),
    src AS (
        SELECT lm.lad_id, lm.metric_id, v.period, v.value_num, v.value_json
        FROM lad_missing lm
        JOIN core.geo_area_link k ON k.parent_id = lm.lad_id
        JOIN core.geo_area c ON c.id = k.child_id AND c.level_code = 'msoa'
        JOIN core.metric_value v ON v.market_code = 'uk' AND v.area_id = c.id AND v.metric_id = lm.metric_id),
    dist AS (
        SELECT lad_id, metric_id, period, jsonb_object_agg(key, total) AS value_json
        FROM (SELECT lad_id, metric_id, period, e.key, sum(e.value::numeric) AS total
              FROM src, jsonb_each_text(src.value_json) e
              WHERE e.value ~ '^-?[0-9.]+$'
              GROUP BY 1, 2, 3, 4) x
        GROUP BY 1, 2, 3)
    INSERT INTO core.metric_value (market_code, area_id, metric_id, period, value_num, value_json, is_imputed, release_id)
    SELECT 'uk', s.lad_id, s.metric_id, s.period, sum(s.value_num), coalesce(d.value_json, '{}'::jsonb), true, %s
    FROM src s LEFT JOIN dist d USING (lad_id, metric_id, period)  -- bảng chỉ có tổng (ts041) không có phân bố
    GROUP BY s.lad_id, s.metric_id, s.period, d.value_json
    ON CONFLICT DO NOTHING"""

CENSUS_LAD_DENSITY_SQL = """
    -- ts006 mật độ (người/km²) của LAD mới = Σ dân số MSOA / Σ diện tích MSOA, diện tích MSOA = dân số / mật độ
    INSERT INTO core.metric_value (market_code, area_id, metric_id, period, value_num, value_json, is_imputed, release_id)
    SELECT 'uk', l.id, 'uk.census21.ts006', d.period, sum(p.value_num) / nullif(sum(p.value_num / nullif(d.value_num, 0)), 0),
           '{}'::jsonb, true, %s
    FROM core.geo_area l
    JOIN core.geo_area_link k ON k.parent_id = l.id
    JOIN core.geo_area c ON c.id = k.child_id AND c.level_code = 'msoa'
    JOIN core.metric_value d ON d.market_code = 'uk' AND d.area_id = c.id AND d.metric_id = 'uk.census21.ts006'
    JOIN core.metric_value p ON p.market_code = 'uk' AND p.area_id = c.id AND p.metric_id = 'uk.census21.ts001'
    WHERE l.market_code = 'uk' AND l.level_code = 'lad' AND l.is_active
      AND NOT EXISTS (SELECT 1 FROM core.metric_value v WHERE v.market_code = 'uk' AND v.area_id = l.id
                      AND v.metric_id = 'uk.census21.ts006')
    GROUP BY l.id, d.period
    ON CONFLICT DO NOTHING"""


def fill_census_lad_from_msoa(ctx: LoadContext) -> int:
    n = ctx.sql("Census: LAD mới = tổng MSOA (is_imputed)", CENSUS_LAD_FILL_SQL, (ctx.release_id,))
    return n + ctx.sql("Census: mật độ LAD mới từ MSOA", CENSUS_LAD_DENSITY_SQL, (ctx.release_id,))


# ===================================================================== UK HPI =====
HPI_METRICS = [  # (cột, id, tên, value_type, unit, decimals, direction)
    ("AveragePrice", "avg_price", "Average house price", "currency", "GBP", 0, 0),
    ("Index", "index", "House price index (Jan 2023 = 100)", "index", None, 1, 0),
    ("12m%Change", "change_12m", "Annual house price change", "percent", "%", 1, 0),
    ("SalesVolume", "sales_volume", "Sales volume", "count", None, 0, 0),
    ("DetachedPrice", "price_detached", "Average price: detached", "currency", "GBP", 0, 0),
    ("SemiDetachedPrice", "price_semi", "Average price: semi-detached", "currency", "GBP", 0, 0),
    ("TerracedPrice", "price_terraced", "Average price: terraced", "currency", "GBP", 0, 0),
    ("FlatPrice", "price_flat", "Average price: flat/maisonette", "currency", "GBP", 0, 0),
]


class HpiLoader(Loader):
    source = "uk_hpi"
    depends_on = ("boundaries_lad", "boundaries_rgn")
    description = "UK HPI full file → metric uk.hpi.* theo LAD/region/nation, hằng tháng"

    def load(self, ctx: LoadContext) -> dict:
        f = ctx.file(r"\.csv$")
        levels = {r[0]: r[1] for r in ctx.conn.execute(
            "SELECT code, level_code FROM core.geo_area WHERE market_code='uk' AND level_code IN ('lad','region','nation')")}
        rows, skipped = [], set()
        with open(f, encoding="utf-8-sig", newline="") as fh:
            rd = csv.DictReader(fh)
            missing = [c for c, *_ in HPI_METRICS if c not in rd.fieldnames]
            if missing:
                raise LoadError(f"UK HPI thiếu cột: {missing}")
            n_src = 0
            for r in rd:
                n_src += 1
                level = levels.get(r["AreaCode"])
                if not level:
                    skipped.add(r["AreaCode"])
                    continue
                period = datetime.strptime(r["Date"], "%d/%m/%Y").date()
                for col, key, *_ in HPI_METRICS:
                    v = num(r[col])
                    if v is not None:
                        rows.append((level, r["AreaCode"], f"uk.hpi.{key}", period, v, None))
        log.info("    %s mã khu trong HPI không có trong geo_area (UK, county, LAD cũ...) → bỏ qua, ví dụ %s",
                 len(skipped), sorted(skipped)[:8])
        for col, key, name, vt, unit, dec, direction in HPI_METRICS:
            register_metric(ctx, f"uk.hpi.{key}", category="housing", dataset="uk.land_registry.hpi", value_type=vt,
                            unit=unit, decimals=dec, direction=direction, period_type="month", native_level="lad",
                            scope="uk.uk", name=name)
        t = mv_stage(ctx, "mv_hpi")
        ctx.copy_rows(t, MV_COLS, rows, "COPY HPI values")
        n, unmatched = insert_values(ctx, t, "UK HPI")
        ctx.drop_stage("mv_hpi")
        return {"rows_in_file": n_src, "values": n, "area_codes_skipped": len(skipped)}


# ======================================================================== PCM =====
PCM_POLLUTANTS = {"pm25": ("PM2.5", r"^mappm25"), "pm10": ("PM10", r"^mappm10"), "no2": ("NO2", r"^mapno2")}


class PcmLoader(Loader):
    source = "defra_pcm"
    depends_on = ("boundaries_lsoa21", "boundaries_msoa21", "boundaries_lad")
    description = ("Defra PCM lưới 1 km → metric uk.pcm.<pollutant> (µg/m³) theo LSOA/MSOA/LAD: "
                   "trung bình các ô có tâm nằm trong khu, khu nhỏ không chứa tâm ô nào → ô gần nhất")

    def load(self, ctx: LoadContext) -> dict:
        year = int(ctx.version)
        period = date(year, 1, 1)
        stats = {}
        for key, (label, pat) in PCM_POLLUTANTS.items():
            f = ctx.file(pat + r".*\.csv$")
            raw = ctx.stage_table(f"pcm_raw_{key}", ["x", "y", "v"], {"x": "float8", "y": "float8", "v": "float8"})
            with open(f, encoding="utf-8-sig", newline="") as fh:
                lines = fh.read().splitlines()
            hdr = next(i for i, l in enumerate(lines) if l.lower().startswith("gridcode"))

            def rows():
                for r in csv.reader(lines[hdr + 1:]):
                    if len(r) >= 4 and r[3].strip() not in ("", "MISSING"):
                        yield (num(r[1]), num(r[2]), num(r[3]))
            n = ctx.copy_rows(raw, ["x", "y", "v"], rows(), f"COPY {f.name}")
            # Tạo bảng điểm bằng CTAS, KHÔNG dùng ALTER + UPDATE: UPDATE tạo HOT chain → index tạo sau đó trong cùng
            # transaction bị đánh dấu indcheckxmin và không dùng được → join không gian chạy hàng chục phút.
            t = f"{STAGE}.pcm_{key}"
            ctx.conn.execute(f"DROP TABLE IF EXISTS {t}")
            ctx.conn.execute(f"""CREATE UNLOGGED TABLE {t} AS
                SELECT v, ST_Transform(ST_SetSRID(ST_MakePoint(x, y), 27700), 4326)::geometry(Point, 4326) AS geom
                FROM {raw}""")
            ctx.conn.execute(f"CREATE INDEX ON {t} USING gist (geom)")
            ctx.conn.execute(f"ANALYZE {t}")
            mid = f"uk.pcm.{key}"
            register_metric(ctx, mid, category="environment", dataset="uk.defra.pcm", value_type="number",
                            unit="µg/m³", decimals=1, direction=-1, period_type="year", native_level="lsoa",
                            aggregation="weighted_mean", scope="uk.uk",
                            name=f"Background {label} concentration (annual mean)")
            s = mv_stage(ctx, "mv_pcm")
            ctx.conn.execute(f"CREATE INDEX ON {s} (level, code)")
            for level in ("lsoa", "msoa", "lad"):
                # 1) mỗi ô lưới → khu chứa tâm ô (dùng GiST index của core.geo_area.geom), gom trung bình theo khu
                ctx.sql(f"PCM {key} → {level} (trung bình ô trong khu)", f"""
                    INSERT INTO {s} (level, code, metric_id, period, value_num)
                    SELECT a.level_code, a.code, %s, %s, avg(p.v)
                    FROM {t} p JOIN core.geo_area a ON a.market_code='uk' AND a.level_code=%s AND a.is_active
                         AND ST_Intersects(a.geom, p.geom)
                    GROUP BY a.level_code, a.code""", (mid, period, level))
                ctx.conn.execute(f"ANALYZE {s}")
                # 2) khu nhỏ không chứa tâm ô nào → ô gần centroid nhất (KNN trên index của bảng ô)
                ctx.sql(f"PCM {key} → {level} (khu không có ô → ô gần nhất)", f"""
                    INSERT INTO {s} (level, code, metric_id, period, value_num)
                    SELECT a.level_code, a.code, %s, %s, n.v
                    FROM core.geo_area a
                    CROSS JOIN LATERAL (SELECT p.v FROM {t} p ORDER BY p.geom <-> a.centroid LIMIT 1) n
                    WHERE a.market_code='uk' AND a.level_code=%s AND a.is_active AND a.geom IS NOT NULL
                      AND NOT EXISTS (SELECT 1 FROM {s} x WHERE x.level = a.level_code AND x.code = a.code)""",
                        (mid, period, level))
            v, _ = insert_values(ctx, s, f"PCM {key}")
            stats[key] = {"grid_cells": n, "values": v}
            ctx.drop_stage("mv_pcm", f"pcm_{key}", f"pcm_raw_{key}")
        return stats



# ======================================================================= PIPR =====
def _slug(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", s.lower()).strip("_")


class PiprLoader(Loader):
    source = "pipr"
    depends_on = ("boundaries_lad", "boundaries_rgn")
    description = "ONS PIPR Table 1 → metric uk.pipr.* (giá thuê, chỉ số, thay đổi năm) theo LAD/region/nation, hằng tháng"

    def load(self, ctx: LoadContext) -> dict:
        import openpyxl
        f = ctx.file(r"\.xlsx$")
        wb = openpyxl.load_workbook(f, read_only=True)
        if "Table 1" not in wb.sheetnames:
            raise LoadError(f"PIPR: không có sheet 'Table 1' (có {wb.sheetnames})")
        it = wb["Table 1"].iter_rows(values_only=True)
        header = None
        for r in it:
            if r and r[0] == "Time period":
                header = [str(c).strip() if c is not None else "" for c in r]
                break
        if not header or "Area code" not in header:
            raise LoadError("PIPR: không tìm thấy dòng header 'Time period'")
        wanted = [c for c in header if c == "Index" or c.startswith("Annual change") or c.startswith("Rental price")]
        idx = {c: header.index(c) for c in wanted}
        log.info("    cột chỉ số: %s", wanted)
        levels = {r[0]: r[1] for r in ctx.conn.execute(
            "SELECT code, level_code FROM core.geo_area WHERE market_code='uk' AND level_code IN ('lad','region','nation')")}
        rows, skipped, n_src = [], set(), 0
        ai = header.index("Area code")
        for r in it:
            if not r or r[0] is None:
                continue
            n_src += 1
            code = r[ai]
            level = levels.get(code)
            if not level:
                skipped.add(code)
                continue
            period = r[0].date() if hasattr(r[0], "date") else datetime.strptime(str(r[0])[:10], "%Y-%m-%d").date()
            for c, i in idx.items():
                v = num(r[i])
                if v is not None:
                    rows.append((level, code, f"uk.pipr.{_slug(c)}", period, v, None))
        log.info("    %s mã khu không có trong geo_area → bỏ qua: %s", len(skipped), sorted(skipped)[:10])
        for c in wanted:
            is_price = c.startswith("Rental price")
            register_metric(ctx, f"uk.pipr.{_slug(c)}", category="rent", dataset="uk.ons.pipr",
                            value_type="currency" if is_price else ("percent" if "change" in c else "index"),
                            unit="GBP/month" if is_price else ("%" if "change" in c else None),
                            decimals=0 if is_price else 1, direction=0, period_type="month", native_level="lad",
                            scope="uk.uk", name=f"Private rent: {c}")
        t = mv_stage(ctx, "mv_pipr")
        ctx.copy_rows(t, MV_COLS, rows, "COPY PIPR values")
        n, unmatched = insert_values(ctx, t, "PIPR")
        ctx.drop_stage("mv_pipr")
        return {"rows_in_sheet": n_src, "values": n, "area_codes_skipped": len(skipped)}


# ===================================================================== Income =====
INCOME_SHEETS = {
    "Total annual income": ("total", "Total annual household income"),
    "Net annual income": ("net", "Net annual household income"),
    "Net income before housing costs": ("net_bhc", "Net household income before housing costs (equivalised)"),
    "Net income after housing costs": ("net_ahc", "Net household income after housing costs (equivalised)"),
}


class IncomeLoader(Loader):
    source = "ons_income_msoa"
    depends_on = ("boundaries_msoa21",)
    description = "ONS small area income (MSOA) → metric uk.income.<loại> (£/năm, có khoảng tin cậy 95%)"

    def load(self, ctx: LoadContext) -> dict:
        import openpyxl
        f = ctx.file(r"\.xlsx$")
        m = re.search(r"(\d{4})$", ctx.version)
        if not m:
            raise LoadError(f"Không đọc được năm tài chính từ version {ctx.version!r}")
        fye = int(m.group(1))
        period = date(fye - 1, 4, 1)  # FYE 2023 = 04/2022–03/2023
        wb = openpyxl.load_workbook(f, read_only=True)
        rows = []
        for sheet, (key, name) in INCOME_SHEETS.items():
            if sheet not in wb.sheetnames:
                raise LoadError(f"Income: thiếu sheet {sheet!r} (có {wb.sheetnames})")
            it = wb[sheet].iter_rows(values_only=True)
            for r in it:
                if r and r[0] == "MSOA code":
                    break
            else:
                raise LoadError(f"Income/{sheet}: không thấy header 'MSOA code'")
            for r in it:
                if r and r[0]:
                    rows.append(("msoa", r[0], f"uk.income.{key}", period, num(r[6]), None, num(r[8]), num(r[7])))
            register_metric(ctx, f"uk.income.{key}", category="economy", dataset="uk.ons.income_msoa",
                            value_type="currency", unit="GBP/year", direction=1, period_type="financial_year",
                            native_level="msoa", aggregation="weighted_mean", scope="uk.ew", name=name,
                            description="Model-based estimate (ONS), mean household income; ci_low/ci_high = 95% CI")
        t = mv_stage(ctx, "mv_income")
        ctx.copy_rows(t, MV_COLS_CI, rows, "COPY income values")
        n, unmatched = insert_values(ctx, t, "income MSOA")
        ctx.drop_stage("mv_income")
        return {"values": n, "unmatched": unmatched, "fye": fye}


class CensusMoreLoader(CensusLoader):
    source = "census2021_more"
    depends_on = ("boundaries_oa21", "onspd", "census2021")
    description = "Census 2021 bảng bổ sung đợt 3 → uk.census21.<ts> (hộ chiếu, di cư, ngôn ngữ, bản sắc: nhạy cảm, không tính điểm)"


class CensusExtraLoader(CensusLoader):
    source = "census2021_extra"
    depends_on = ("boundaries_oa21", "onspd", "census2021")
    description = "Census 2021 bảng bổ sung (nơi sinh, thời gian cư trú, tiếng Anh, khuyết tật, chăm sóc, nghề nghiệp) → uk.census21.<ts>"
