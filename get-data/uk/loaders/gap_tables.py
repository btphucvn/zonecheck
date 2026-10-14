"""Nguồn dạng bảng bổ sung để lấp khoảng trống so với đối thủ (xem DATA_GAP_TRACKER.md).

Mỗi loader nạp trong một transaction, có kiểm tra số dòng/khoảng giá trị/đối chiếu chéo; lỗi → rollback.
"""
from __future__ import annotations

import csv
import io
import json
import logging
import re
import zipfile
from datetime import date

from .base import LoadContext, LoadError, Loader, zip_members
from .metrics import register_metric

log = logging.getLogger(__name__)


def _ddl(ctx: LoadContext, ddl: str) -> None:
    for stmt in ddl.strip().split(";\n"):
        if stmt.strip():
            ctx.conn.execute(stmt)


def _num(v):
    if v is None:
        return None
    s = str(v).strip().replace(",", "")
    if s in ("", "[x]", "[z]", "[c]", "[u]", "x", "z", "c", "..", "-", "n/a", "None"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _set_subject_poi(ctx: LoadContext, ids: list[str]) -> None:
    ctx.conn.execute("UPDATE core.metric SET subject = 'poi' WHERE id = ANY(%s)", (ids,))


# ------------------------------------------------------------------ RoFRS -----
ROFRS_DDL = """
CREATE TABLE IF NOT EXISTS mkt_uk.postcode_rofrs (
    postcode text PRIMARY KEY,
    properties int, residential int,
    res_very_low int, res_low int, res_medium int, res_high int,
    all_very_low int, all_low int, all_medium int, all_high int,
    release_id bigint REFERENCES ops.dataset_release (id)
);
COMMENT ON TABLE mkt_uk.postcode_rofrs IS 'EA NaFRA2 Risk of Flooding from Rivers and Sea: số nhà/cơ sở theo mức rủi ro trong từng postcode có rủi ro (England). Postcode không có trong bảng = không có nhà nào trong vùng rủi ro.'"""


class RofrsLoader(Loader):
    source = "ea_rofrs_postcodes"
    depends_on = ("onspd", "census2021")
    description = "EA RoFRS theo postcode → mkt_uk.postcode_rofrs + uk.zone.rs_{high,medium_plus,any}_share theo LSOA"

    def load(self, ctx: LoadContext) -> dict:
        z = ctx.file(r"RoFRS_Postcodes_in_Areas_at_Risk\.zip$")
        members = zip_members(z, r"Postcodes_AtRisk\.csv$")
        if len(members) != 1:
            raise LoadError(f"Cần 1 file Postcodes_AtRisk.csv, thấy {members}")
        cols = ["PC", "cntpc", "RES_cntpc", "NRP_cntpc", "UNC_cntpc"] + [
            f"{t}_CNT_{b}" for b in ("VeryLow", "Low", "Medium", "High") for t in ("RES", "NRP", "UNC", "TOT")]
        with zipfile.ZipFile(z) as zf, zf.open(members[0]) as f:
            rd = csv.DictReader(io.TextIOWrapper(f, encoding="utf-8-sig", newline=""))
            miss = [c for c in cols if c not in rd.fieldnames]
            if miss:
                raise LoadError(f"RoFRS thiếu cột {miss}")
            rows = [[r["PC"].strip().upper()] + [int(float(r[c] or 0)) for c in cols[1:]] for r in rd]
        st = ctx.stage_table("rofrs", ["pc"] + [c.lower() for c in cols[1:]], {c.lower(): "int" for c in cols[1:]})
        n = ctx.copy_rows(st, ["pc"] + [c.lower() for c in cols[1:]], rows, "COPY RoFRS postcodes")
        if not 100_000 <= n <= 400_000:
            raise LoadError(f"Số postcode có rủi ro bất thường: {n:,}")
        _ddl(ctx, ROFRS_DDL)
        ctx.sql("TRUNCATE postcode_rofrs", "TRUNCATE mkt_uk.postcode_rofrs")
        full_re = "^[A-Z]{1,2}[0-9][A-Z0-9]? [0-9][A-Z]{2}$"
        n_full = ctx.scalar(f"SELECT count(*) FROM {st} WHERE pc ~ %s", (full_re,))
        n_wild = ctx.scalar(f"SELECT count(*) FROM {st} WHERE pc ~ '^[A-Z]{{1,2}}[0-9][A-Z0-9]? [0-9][A-Z]?\\*$'")
        log.info("    %s postcode đầy đủ, %s dòng gộp dạng 'AB1 2C*' (ẩn danh), %s mã khác", f"{n_full:,}", f"{n_wild:,}",
                 f"{n - n_full - n_wild:,}")
        ins = ctx.sql("INSERT postcode_rofrs (postcode đầy đủ)", f"""
            INSERT INTO mkt_uk.postcode_rofrs SELECT p.postcode, s.cntpc, s.res_cntpc,
                s.res_cnt_verylow, s.res_cnt_low, s.res_cnt_medium, s.res_cnt_high,
                s.tot_cnt_verylow, s.tot_cnt_low, s.tot_cnt_medium, s.tot_cnt_high, %s
            FROM {st} s JOIN mkt_uk.postcode p ON p.postcode_compact = replace(s.pc, ' ', '')
            WHERE s.pc ~ %s""", (ctx.release_id, full_re))
        match = ins / max(n_full, 1)
        log.info("    %s/%s postcode đầy đủ khớp ONSPD (%.1f%%)", f"{ins:,}", f"{n_full:,}", match * 100)
        if match < 0.97:
            raise LoadError(f"Chỉ {match:.1%} postcode RoFRS khớp ONSPD")
        # dòng gộp: chia đều số nhà cho các postcode đang dùng có cùng tiền tố (sector + chữ đầu của unit)
        ctx.sql("dòng gộp → postcode cùng tiền tố", f"""
            CREATE TEMP TABLE rofrs_wild ON COMMIT DROP AS
            SELECT p.lsoa21_code, w.pc,
                   s.res_cnt_verylow::float8 / count(*) OVER (PARTITION BY w.pc) AS vl,
                   s.res_cnt_low::float8 / count(*) OVER (PARTITION BY w.pc) AS lo,
                   s.res_cnt_medium::float8 / count(*) OVER (PARTITION BY w.pc) AS me,
                   s.res_cnt_high::float8 / count(*) OVER (PARTITION BY w.pc) AS hi
            FROM (SELECT pc, rtrim(pc, '*') AS pre FROM {st} WHERE pc LIKE '%%*') w
            JOIN {st} s USING (pc)
            JOIN mkt_uk.postcode p ON p.sector = split_part(w.pre, ' ', 1) || ' ' || left(split_part(w.pre, ' ', 2), 1)
                 AND p.status = 'live' AND p.postcode LIKE w.pre || '%%'""")
        wild_ok = ctx.scalar("SELECT count(DISTINCT pc) FROM rofrs_wild")
        log.info("    %s/%s dòng gộp tìm được postcode cùng tiền tố", f"{wild_ok:,}", f"{n_wild:,}")
        if n_wild and wild_ok < 0.9 * n_wild:
            raise LoadError(f"Chỉ {wild_ok}/{n_wild} dòng gộp khớp postcode")
        ctx.sql("tổng theo LSOA", """
            CREATE TEMP TABLE rofrs_lsoa ON COMMIT DROP AS
            SELECT lsoa21_code, sum(vl) vl, sum(lo) lo, sum(me) me, sum(hi) hi FROM (
                SELECT p.lsoa21_code, r.res_very_low vl, r.res_low lo, r.res_medium me, r.res_high hi
                FROM mkt_uk.postcode_rofrs r JOIN mkt_uk.postcode p USING (postcode)
                UNION ALL SELECT lsoa21_code, vl, lo, me, hi FROM rofrs_wild) x GROUP BY 1""")
        tot = ctx.conn.execute("SELECT sum(vl + lo + me + hi)::bigint, sum(hi)::bigint, sum(me)::bigint FROM rofrs_lsoa").fetchone()
        log.info("    nhà ở có rủi ro sông/biển: %s (cao %s, trung bình %s)", *[f"{x:,}" for x in tot])
        # NaFRA2 (2025): ~2,4 triệu nhà/cơ sở England có rủi ro sông/biển (mọi mức), phần lớn ở mức thấp/rất thấp
        if not (800_000 <= tot[0] <= 4_000_000 and tot[1] < tot[0] * 0.5):
            raise LoadError(f"Tổng số nhà rủi ro sông/biển bất thường: {tot}")

        period = date.fromisoformat(ctx.version)
        defs = {"rs_high": ("Homes at high risk of flooding from rivers or the sea (3.3%+ a year)", "hi", True),
                "rs_medium_plus": ("Homes at medium or high risk of flooding from rivers or the sea (1%+ a year)",
                                   "me + hi", False),
                "rs_any": ("Homes at any risk of flooding from rivers or the sea", "vl + lo + me + hi", False)}
        for k, (name, _, scorable) in defs.items():
            register_metric(ctx, f"uk.zone.{k}_share", category="risk", dataset="uk.ea.rofrs_postcodes",
                            value_type="percent", unit="%", decimals=1, direction=-1, period_type="snapshot",
                            native_level="lsoa", aggregation="weighted_mean", scope="uk.england", name=name,
                            scorable=scorable, description="EA NaFRA2 RoFRS residential properties at risk ÷ households (Census 2021)")
        ids = [f"uk.zone.{k}_share" for k in defs]
        ctx.sql("xoá giá trị cũ", "DELETE FROM core.metric_value WHERE market_code='uk' AND metric_id = ANY(%s)", (ids,))
        n_v = 0
        for k, (_, expr, _) in defs.items():
            n_v += ctx.sql(f"uk.zone.{k}_share", f"""
                INSERT INTO core.metric_value (market_code, area_id, metric_id, period, value_num, sample_size, release_id)
                SELECT 'uk', a.id, %s, %s, least(100, round((100.0 * coalesce(r.n, 0) / hh.value_num)::numeric, 2)),
                       coalesce(r.n, 0), %s
                FROM core.geo_area a
                JOIN core.metric_value hh ON hh.market_code='uk' AND hh.area_id = a.id AND hh.metric_id = 'uk.census21.ts041'
                LEFT JOIN (SELECT lsoa21_code, {expr} n FROM rofrs_lsoa) r ON r.lsoa21_code = a.code
                WHERE a.market_code='uk' AND a.level_code='lsoa' AND a.code LIKE 'E%%' AND hh.value_num > 0""",
                           (f"uk.zone.{k}_share", period, ctx.release_id))
        ctx.drop_stage("rofrs")
        return {"postcodes": ins, "residential_at_risk": tot[0], "values": n_v}


# ---------------------------------------------------------- affordability -----
class AffordabilityLoader(Loader):
    source = "ons_affordability"
    depends_on = ("boundaries_lad",)
    description = "ONS tỷ lệ giá nhà/thu nhập nơi làm việc (trung vị 5c, nhóm thấp 6c) theo LAD, mọi năm → uk.afford.*"

    def load(self, ctx: LoadContext) -> dict:
        import openpyxl
        wb = openpyxl.load_workbook(ctx.file(r"\.xlsx$"), read_only=True, data_only=True)
        sheets = {"5c": ("uk.afford.median_ratio", "House price to earnings ratio (median)"),
                  "6c": ("uk.afford.lq_ratio", "House price to earnings ratio (lower quartile)")}
        rows = []
        for sh, (mid, _) in sheets.items():
            it = wb[sh].iter_rows(values_only=True)
            header = None
            for r in it:
                if r and r[0] == "Country/Region code":
                    header = list(r)
                    break
            if not header or header[2] != "Local authority code":
                raise LoadError(f"Sheet {sh}: không thấy dòng tiêu đề LA (có {header})")
            years = [(i, int(str(h)[:4])) for i, h in enumerate(header) if h and re.fullmatch(r"\d{4}", str(h))]
            for r in it:
                if not r or not r[2] or not re.match(r"^[EW]0\d{7}$", str(r[2])):
                    continue
                for i, y in years:
                    v = _num(r[i])
                    if v is not None:
                        rows.append(("lad", r[2], mid, date(y, 9, 30), v))
        if len(rows) < 8000:
            raise LoadError(f"Chỉ có {len(rows)} giá trị tỷ lệ (kỳ vọng ~300 LA × ~28 năm × 2)")
        st = ctx.stage_table("aff", ["level", "code", "metric_id", "period", "value_num"],
                             {"period": "date", "value_num": "float8"})
        ctx.copy_rows(st, ["level", "code", "metric_id", "period", "value_num"], rows, "COPY tỷ lệ")
        ctx.check_range("tỷ lệ", st, "value_num", 0.5, 60)  # năm 1997 có LA < 1
        latest = ctx.scalar(f"SELECT max(period) FROM {st}")
        eng = ctx.scalar(f"SELECT percentile_cont(0.5) WITHIN GROUP (ORDER BY value_num) FROM {st} WHERE metric_id='uk.afford.median_ratio' AND period=%s", (latest,))
        kc = ctx.scalar(f"SELECT value_num FROM {st} WHERE code='E09000020' AND metric_id='uk.afford.median_ratio' AND period=%s", (latest,))
        log.info("    năm mới nhất %s: trung vị các LA %.2f, Kensington & Chelsea %.2f", latest, eng, kc or 0)
        if kc is None or kc < 2 * eng:  # đối chiếu định hướng: K&C luôn đắt nhất nước
            raise LoadError(f"Kiểm tra định hướng thất bại: K&C {kc} không gấp đôi trung vị {eng}")
        for mid, name in sheets.values():
            register_metric(ctx, mid, category="housing", dataset="uk.ons.affordability", value_type="number",
                            unit="ratio", decimals=2, direction=-1, period_type="year", native_level="lad",
                            aggregation="none", scope="uk.ew", name=name, scorable=False,
                            description="ONS: median (or lower quartile) house price ÷ workplace-based annual earnings, year to September")
        ctx.sql("xoá giá trị cũ", "DELETE FROM core.metric_value WHERE market_code='uk' AND metric_id LIKE 'uk.afford.%%'")
        n = ctx.sql("INSERT tỷ lệ", f"""
            INSERT INTO core.metric_value (market_code, area_id, metric_id, period, value_num, release_id)
            SELECT 'uk', a.id, s.metric_id, s.period, s.value_num, %s FROM {st} s
            JOIN core.geo_area a ON a.market_code='uk' AND a.level_code='lad' AND a.code = s.code""", (ctx.release_id,))
        if n < 0.8 * len(rows):
            raise LoadError(f"Chỉ {n}/{len(rows)} giá trị khớp mã LAD hiện hành")
        ctx.drop_stage("aff")
        return {"values": n, "latest": str(latest)}


# ------------------------------------------------------------------- PTAL -----
class PtalLoader(Loader):
    source = "london_ptal"
    depends_on = ("boundaries_lsoa21",)
    description = "London PTAL 2015 → uk.london.ptai (LSOA) + mkt_uk.ptal_grid (lưới 100 m)"

    def load(self, ctx: LoadContext) -> dict:
        rows = []
        with open(ctx.file(r"^LSOA2011_AvPTAI2015\.csv$"), encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                v = _num(r["AvPTAI2015"])
                if v is not None:
                    rows.append(("lsoa", r["LSOA2011"], "uk.london.ptai", date(2015, 12, 31), v,
                                 json.dumps({"ptal": r["PTAL"]})))
        if not 4500 <= len(rows) <= 5000:
            raise LoadError(f"Số LSOA London bất thường: {len(rows)} (London có 4.835 LSOA 2011)")
        st = ctx.stage_table("ptal", ["level", "code", "metric_id", "period", "value_num", "value_json"],
                             {"period": "date", "value_num": "float8", "value_json": "jsonb"})
        ctx.copy_rows(st, ["level", "code", "metric_id", "period", "value_num", "value_json"], rows, "COPY PTAL")
        ctx.check_range("PTAI", st, "value_num", 0, 500)
        register_metric(ctx, "uk.london.ptai", category="transport", dataset="uk.gla.ptal", value_type="index",
                        decimals=1, direction=1, period_type="snapshot", native_level="lsoa",
                        aggregation="weighted_mean", scope="uk.england", name="Public transport accessibility (PTAI, London)",
                        scorable=False, description="TfL PTAL 2015, mean Public Transport Accessibility Index; value_json.ptal = band 0–6b")
        ctx.sql("xoá giá trị cũ", "DELETE FROM core.metric_value WHERE market_code='uk' AND metric_id='uk.london.ptai'")
        n = ctx.sql("INSERT PTAI", f"""
            INSERT INTO core.metric_value (market_code, area_id, metric_id, period, value_num, value_json, release_id)
            SELECT 'uk', a.id, s.metric_id, s.period, s.value_num, s.value_json, %s FROM {st} s
            JOIN core.geo_area a ON a.market_code='uk' AND a.level_code='lsoa' AND a.code = s.code""", (ctx.release_id,))
        if n < 0.85 * len(rows):  # LSOA 2011→2021 đổi mã ~5%
            raise LoadError(f"Chỉ {n}/{len(rows)} LSOA khớp mã 2021")
        city = ctx.scalar(f"SELECT avg(value_num) FROM {st} WHERE code LIKE 'E0100000%%'")  # City of London
        if city is None or city < 40:
            raise LoadError(f"PTAI City of London bất thường: {city}")
        ctx.drop_stage("ptal")
        return {"values": n, "lsoa_2011": len(rows), **self._grid(ctx)}

    def _grid(self, ctx: LoadContext) -> dict:
        """Lưới 100 m (tâm ô, BNG) → mkt_uk.ptal_grid để tra PTAL đúng tại một địa chỉ (ô gần nhất)."""
        import openpyxl
        z = ctx.file(r"Grid.*Values.*\.zip$")
        with zipfile.ZipFile(z) as zf:
            m = [n for n in zf.namelist() if n.lower().endswith(".xlsx")]
            if len(m) != 1:
                raise LoadError(f"Cần 1 file xlsx trong {z.name}, thấy {m}")
            wb = openpyxl.load_workbook(io.BytesIO(zf.read(m[0])), read_only=True)
        it = wb.worksheets[0].iter_rows(values_only=True)
        head = [str(h) for h in next(it)]
        need = ["X", "Y", "AI2015", "PTAL2015"]
        if any(c not in head for c in need):
            raise LoadError(f"Lưới PTAL thiếu cột {need}: {head}")
        ix = [head.index(c) for c in need]
        bands = {"0", "1a", "1b", "2", "3", "4", "5", "6a", "6b"}
        rows, odd = [], 0
        for r in it:
            x, y, ai, ptal = (r[i] for i in ix)
            ptal = str(ptal).strip()
            if x is None or y is None or ptal not in bands:
                odd += 1
                continue
            rows.append((float(x), float(y), float(ai or 0), ptal))
        wb.close()
        if not 150_000 <= len(rows) <= 170_000 or odd > 100:
            raise LoadError(f"Lưới PTAL: {len(rows)} ô hợp lệ, {odd} ô lạ")
        ctx.conn.execute("""CREATE TABLE IF NOT EXISTS mkt_uk.ptal_grid (
            id serial PRIMARY KEY, ai real NOT NULL, ptal text NOT NULL, geom geometry(Point, 4326) NOT NULL,
            release_id bigint REFERENCES ops.dataset_release (id))""")
        ctx.conn.execute("CREATE INDEX IF NOT EXISTS ptal_grid_geom ON mkt_uk.ptal_grid USING gist (geom)")
        ctx.conn.execute("COMMENT ON TABLE mkt_uk.ptal_grid IS 'TfL PTAL 2015, lưới 100 m (tâm ô) London: chỉ số AI và dải PTAL 0–6b. London Datastore, OGL.'")
        st = ctx.stage_table("ptal_grid", ["x", "y", "ai", "ptal"], {"x": "float8", "y": "float8", "ai": "float8"})
        ctx.copy_rows(st, ["x", "y", "ai", "ptal"], rows, "COPY lưới PTAL")
        ctx.sql("thay lưới PTAL", "TRUNCATE mkt_uk.ptal_grid")
        n = ctx.sql("ghi lưới PTAL", f"""INSERT INTO mkt_uk.ptal_grid (ai, ptal, geom, release_id)
            SELECT ai, ptal, ST_Transform(ST_SetSRID(ST_MakePoint(x, y), 27700), 4326), %s FROM {st}""", (ctx.release_id,))
        # kiểm tra định hướng: ô gần Bank (trung tâm City) phải là 6b
        bank = ctx.scalar("""SELECT ptal FROM mkt_uk.ptal_grid ORDER BY geom <-> ST_SetSRID(ST_MakePoint(-0.0886, 51.5133), 4326) LIMIT 1""")
        log.info("    lưới PTAL: %s ô, ô gần Bank = %s", f"{n:,}", bank)
        if bank != "6b":
            raise LoadError(f"PTAL tại Bank phải là 6b, thấy {bank}")
        ctx.conn.execute("ANALYZE mkt_uk.ptal_grid")
        ctx.drop_stage("ptal_grid")
        return {"grid_cells": n}


# ------------------------------------------------------------ fly-tipping -----
class FlyTippingLoader(Loader):
    source = "fly_tipping"
    depends_on = ("boundaries_lad", "census2021")
    description = "Defra fly-tipping: số vụ theo LA × năm tài chính → uk.flytip.incidents (+ per_1k)"

    def load(self, ctx: LoadContext) -> dict:
        from pathlib import Path

        from common.ods import iter_ods_rows
        path = ctx.file(r"reported_by_LAs.*\.ods$")
        header, rows = None, []
        for r in iter_ods_rows(Path(path), "LA_incidents"):
            if header is None:
                if r and r[0] == "Year" and "ONS Code" in r:
                    header = r
                    i_year, i_code, i_tot = 0, r.index("ONS Code"), r.index("Total Incidents")
                continue
            if len(r) <= i_tot or not r[i_code] or not re.match(r"^E0\d{7}$", str(r[i_code])):
                continue
            m = re.match(r"^(\d{4})-(\d{2})$", str(r[i_year]).strip())
            v = _num(r[i_tot])
            if m and v is not None:
                rows.append(("lad", r[i_code], "uk.flytip.incidents", date(int(m.group(1)) + 1, 3, 31), v))
        if header is None:
            raise LoadError("Không thấy dòng tiêu đề Year/ONS Code/Total Incidents trong sheet LA_incidents")
        if len(rows) < 3000:
            raise LoadError(f"Chỉ có {len(rows)} giá trị (kỳ vọng ~300 LA × 13 năm)")
        st = ctx.stage_table("flytip", ["level", "code", "metric_id", "period", "value_num"],
                             {"period": "date", "value_num": "float8"})
        ctx.copy_rows(st, ["level", "code", "metric_id", "period", "value_num"], rows, "COPY fly-tipping")
        ctx.check_range("số vụ", st, "value_num", 0, 200_000)
        latest = ctx.scalar(f"SELECT max(period) FROM {st}")
        total = ctx.scalar(f"SELECT sum(value_num) FROM {st} WHERE period = %s", (latest,))
        log.info("    năm %s: tổng %s vụ England", latest, f"{total:,.0f}")
        if not 700_000 <= total <= 1_500_000:  # England ~1,1 triệu vụ/năm
            raise LoadError(f"Tổng số vụ fly-tipping bất thường: {total:,.0f}")
        register_metric(ctx, "uk.flytip.incidents", category="environment", dataset="uk.defra.fly_tipping",
                        value_type="count", decimals=0, direction=-1, period_type="financial_year", native_level="lad",
                        aggregation="sum", scope="uk.england", name="Fly-tipping incidents reported by the council",
                        scorable=False, description="Defra WasteDataFlow, incidents dealt with by the local authority")
        register_metric(ctx, "uk.flytip.per_1k", category="environment", dataset="uk.defra.fly_tipping",
                        value_type="rate", unit="per 1,000", decimals=1, direction=-1, period_type="financial_year",
                        native_level="lad", aggregation="none", scope="uk.england",
                        name="Fly-tipping incidents per 1,000 residents", scorable=False,
                        description="Incidents ÷ Census 2021 residents × 1,000")
        ctx.sql("xoá giá trị cũ", "DELETE FROM core.metric_value WHERE market_code='uk' AND metric_id LIKE 'uk.flytip.%%'")
        n = ctx.sql("INSERT số vụ", f"""
            INSERT INTO core.metric_value (market_code, area_id, metric_id, period, value_num, release_id)
            SELECT 'uk', a.id, s.metric_id, s.period, s.value_num, %s FROM {st} s
            JOIN core.geo_area a ON a.market_code='uk' AND a.level_code='lad' AND a.code = s.code""", (ctx.release_id,))
        n2 = ctx.sql("INSERT per 1.000 dân", f"""
            INSERT INTO core.metric_value (market_code, area_id, metric_id, period, value_num, release_id)
            SELECT 'uk', a.id, 'uk.flytip.per_1k', s.period, round((1000.0 * s.value_num / p.value_num)::numeric, 2), %s
            FROM {st} s JOIN core.geo_area a ON a.market_code='uk' AND a.level_code='lad' AND a.code = s.code
            JOIN core.metric_value p ON p.market_code='uk' AND p.area_id = a.id AND p.metric_id='uk.census21.ts001'
            WHERE p.value_num > 0""", (ctx.release_id,))
        ctx.drop_stage("flytip")
        return {"values": n, "per_1k": n2, "latest_total": total}


# -------------------------------------------------------------------- LHA -----
class LhaLoader(Loader):
    source = "lha_rates"
    description = "VOA Local Housing Allowance (bảng LHA hiện hành) → mkt_uk.lha_rate (BRMA × loại phòng, £/tuần)"
    DDL = """
CREATE TABLE IF NOT EXISTS mkt_uk.lha_rate (
    brma text NOT NULL, category char(1) NOT NULL, weekly_gbp numeric(8,2) NOT NULL, valid_from date NOT NULL,
    release_id bigint REFERENCES ops.dataset_release (id), PRIMARY KEY (brma, category, valid_from)
);
COMMENT ON TABLE mkt_uk.lha_rate IS 'VOA Local Housing Allowance (England): trần trợ cấp thuê theo Broad Rental Market Area. A = phòng chung, B = 1 phòng ngủ, C = 2, D = 3, E = 4.'"""

    def load(self, ctx: LoadContext) -> dict:
        import openpyxl
        wb = openpyxl.load_workbook(ctx.file(r"LHA_TABLES\.xlsx$"), read_only=True, data_only=True)
        best = None
        for ws in wb.worksheets:
            first = next(ws.iter_rows(values_only=True, max_row=1), ())
            m = re.search(r"LHA April (\d{4})", " ".join(str(c) for c in first if c))
            if m and (best is None or int(m.group(1)) > best[0]):
                best = (int(m.group(1)), ws)
        if not best:
            raise LoadError("Không thấy sheet 'LHA April YYYY'")
        year, ws = best
        rows = []
        for r in ws.iter_rows(values_only=True, min_row=3):
            if not r or not r[0] or r[0] == "BRMA":
                continue
            for cat, v in zip("ABCDE", r[1:6]):
                v = _num(v)
                if v is not None:
                    rows.append((str(r[0]).strip(), cat, v, date(year, 4, 1), ctx.release_id))
        if not 140 * 5 <= len(rows) <= 160 * 5:
            raise LoadError(f"Số dòng LHA bất thường: {len(rows)} (England có 152 BRMA × 5 loại)")
        _ddl(ctx, self.DDL)
        ctx.sql("xoá cùng năm", "DELETE FROM mkt_uk.lha_rate WHERE valid_from = %s", (date(year, 4, 1),))
        n = ctx.copy_rows("mkt_uk.lha_rate", ["brma", "category", "weekly_gbp", "valid_from", "release_id"], rows, "COPY LHA")
        ctx.check_range("£/tuần", "mkt_uk.lha_rate", "weekly_gbp", 40, 1000, where="valid_from = %s" % f"'{date(year, 4, 1)}'")
        return {"rows": n, "year": year}


# --------------------------------------------------------------- tribunal -----
class TribunalLoader(Loader):
    source = "tribunal_decisions"
    depends_on = ("onspd",)
    description = "gov.uk Residential Property Tribunal decisions → mkt_uk.tribunal_decision (postcode tách từ tiêu đề)"
    DDL = """
CREATE TABLE IF NOT EXISTS mkt_uk.tribunal_decision (
    link text PRIMARY KEY, title text NOT NULL, case_ref text, case_type text, postcode text, published_at timestamptz,
    description text, release_id bigint REFERENCES ops.dataset_release (id)
);
CREATE INDEX IF NOT EXISTS tribunal_decision_postcode ON mkt_uk.tribunal_decision (postcode);
COMMENT ON TABLE mkt_uk.tribunal_decision IS 'Quyết định First-tier Tribunal (Property Chamber) công bố trên gov.uk. case_type = mã loại trong số vụ (LSC = phí dịch vụ, HMF = lệnh hoàn tiền thuê, LIS, LDC…). Chỉ hiển thị tổng hợp theo khu.'"""
    PC = re.compile(r"\b([A-Z]{1,2}\d[A-Z\d]?)\s*(\d[A-Z]{2})\b")
    REF = re.compile(r"\b([A-Z]{3}/[0-9A-Z]{4}/([A-Z]{2,4})/\d{4}/\d{4})\b")

    def load(self, ctx: LoadContext) -> dict:
        data = json.loads(ctx.file(r"rpt_decisions\.json$").read_text())
        rows, seen = [], set()
        for d in data:
            if d["link"] in seen:
                continue
            seen.add(d["link"])
            t = d.get("title") or ""
            pcs = self.PC.findall(t.upper())
            ref = self.REF.search(t)
            rows.append((d["link"], t, ref.group(1) if ref else None, ref.group(2) if ref else None,
                         f"{pcs[-1][0]} {pcs[-1][1]}" if pcs else None, d.get("public_timestamp"),
                         d.get("description"), ctx.release_id))
        if len(rows) < 5000:
            raise LoadError(f"Chỉ có {len(rows)} quyết định")
        _ddl(ctx, self.DDL)
        ctx.sql("TRUNCATE tribunal_decision", "TRUNCATE mkt_uk.tribunal_decision")
        n = ctx.copy_rows("mkt_uk.tribunal_decision", ["link", "title", "case_ref", "case_type", "postcode",
                                                       "published_at", "description", "release_id"], rows, "COPY tribunal")
        ok = ctx.scalar("""SELECT count(*) FROM mkt_uk.tribunal_decision t JOIN mkt_uk.postcode p ON p.postcode = t.postcode""")
        log.info("    %s/%s quyết định có postcode khớp ONSPD", f"{ok:,}", f"{n:,}")
        if ok < 0.6 * n:
            raise LoadError(f"Chỉ {ok}/{n} quyết định tách được postcode hợp lệ")
        ctx.sql("postcode không khớp → NULL", """UPDATE mkt_uk.tribunal_decision t SET postcode = NULL
            WHERE postcode IS NOT NULL AND NOT EXISTS (SELECT 1 FROM mkt_uk.postcode p WHERE p.postcode = t.postcode)""")
        return {"decisions": n, "with_postcode": ok}


# --------------------------------------------------------------------- MP -----
# Tên khu bầu cử khác nhau giữa OS Boundary-Line và Parliament: OS có hậu tố "Boro/Co/Burgh Const", "St." thay "St",
# không dấu tiếng Wales ("Ynys Mon" ↔ "Ynys Môn"), "&" ↔ "and".
MP_NORM = ("lower(regexp_replace(regexp_replace(public.unaccent(replace(regexp_replace({e}, '\\s+(Boro|Co|Burgh)\\s+Const$', ''), "
           "'&', 'and')), '[.,]', '', 'g'), '\\s+', ' ', 'g'))")


class MpsLoader(Loader):
    source = "parliament_mps"
    depends_on = ("os_boundary_line",)
    description = "UK Parliament: MP đương nhiệm → mkt_uk.mp (theo tên khu bầu cử)"
    DDL = """
CREATE TABLE IF NOT EXISTS mkt_uk.mp (
    member_id int PRIMARY KEY, name text NOT NULL, party text, constituency text NOT NULL,
    constituency_id int, member_since date, thumbnail_url text,
    release_id bigint REFERENCES ops.dataset_release (id)
);
COMMENT ON TABLE mkt_uk.mp IS 'Nghị sĩ (House of Commons) đương nhiệm theo khu bầu cử (UK Parliament Members API, Open Parliament Licence). Ghép với khu bầu cử (core.zone admin.westminster_const) qua gss_code.';
ALTER TABLE mkt_uk.mp ADD COLUMN IF NOT EXISTS gss_code text;
COMMENT ON COLUMN mkt_uk.mp.gss_code IS 'Mã GSS khu bầu cử (= core.zone.external_id của admin.westminster_const), ghép theo tên đã chuẩn hoá. NULL ở Bắc Ireland (Boundary-Line không có).';
CREATE INDEX IF NOT EXISTS mp_gss_code_idx ON mkt_uk.mp (gss_code)"""

    def load(self, ctx: LoadContext) -> dict:
        data = json.loads(ctx.file(r"mps\.json$").read_text())
        rows = []
        for m in data:
            hm = m.get("latestHouseMembership") or {}
            since = (hm.get("membershipStartDate") or "")[:10] or None
            rows.append((m["id"], m.get("nameDisplayAs"), (m.get("latestParty") or {}).get("name"),
                         hm.get("membershipFrom"), hm.get("membershipFromId"), since, m.get("thumbnailUrl"), ctx.release_id))
        if not 600 <= len(rows) <= 660 or any(not r[3] for r in rows):
            raise LoadError(f"Danh sách MP bất thường: {len(rows)} dòng, thiếu khu bầu cử: {sum(1 for r in rows if not r[3])}")
        if len({r[3] for r in rows}) < len(rows) - 5:
            raise LoadError("Trùng khu bầu cử bất thường")
        _ddl(ctx, self.DDL)
        ctx.sql("TRUNCATE mp", "TRUNCATE mkt_uk.mp")
        n = ctx.copy_rows("mkt_uk.mp", ["member_id", "name", "party", "constituency", "constituency_id", "member_since",
                                        "thumbnail_url", "release_id"], rows, "COPY MP")
        matched = ctx.sql("gắn mã GSS khu bầu cử", f"""
            UPDATE mkt_uk.mp m SET gss_code = z.external_id
            FROM core.zone z WHERE z.market_code = 'uk' AND z.type_code = 'admin.westminster_const'
              AND {MP_NORM.format(e='z.name')} = {MP_NORM.format(e='m.constituency')}""")
        n_const = ctx.scalar("SELECT count(*) FROM core.zone WHERE market_code='uk' AND type_code='admin.westminster_const'")
        dup = ctx.scalar("SELECT count(*) - count(DISTINCT gss_code) FROM mkt_uk.mp WHERE gss_code IS NOT NULL")
        unmatched = [r[0] for r in ctx.conn.execute(f"""SELECT z.name FROM core.zone z
            WHERE z.market_code='uk' AND z.type_code='admin.westminster_const'
              AND NOT EXISTS (SELECT 1 FROM mkt_uk.mp m WHERE m.gss_code = z.external_id) LIMIT 10""")]
        log.info("    MP gắn được khu bầu cử: %s/%s khu GB (trùng %s), khu chưa có MP: %s", matched, n_const, dup, unmatched)
        if dup or matched < n_const - 3:  # có thể trống ghế vài khu (MP từ chức, chờ bầu bổ sung)
            raise LoadError(f"Ghép MP ↔ khu bầu cử kém: {matched}/{n_const}, trùng {dup}, chưa khớp {unmatched}")
        return {"mps": n, "constituencies_matched": matched}


# --------------------------------------------------------------- nhà trẻ -----
CHILDCARE_FIELDS = ["Provider URN", "Provider type", "Provider subtype", "Provider status", "Provider name", "Postcode",
                    "Places", "EYR OEIF/CIF: Most recent full: Overall effectiveness",
                    "EYR OEIF/CIF: Most recent full: Inspection date", "EYR REIF: Most recent: Inspection date",
                    "EYR REIF: Most recent: Curriculum and teaching", "EYR REIF: Most recent: Children's welfare and wellbeing"]


class ChildcareLoader(Loader):
    source = "ofsted_childcare"
    depends_on = ("onspd",)
    description = "Ofsted childcare (nhà trẻ/nhóm trẻ có địa chỉ công khai) → core.poi childcare.nursery + poi_rating ofsted_ey"

    def load(self, ctx: LoadContext) -> dict:
        from pathlib import Path

        from common.ods import iter_ods_rows
        path = ctx.file(r"provider_level.*\.ods$")
        header, rows = None, []
        for r in iter_ods_rows(Path(path), "D1_Childcare_providers"):
            if header is None:
                if r and r[0] == "Web link" and "Provider URN" in r:
                    header = r
                    ix = {h: i for i, h in enumerate(header)}
                    need = ["Provider URN", "Provider type", "Provider subtype", "Provider status", "Provider name",
                            "Postcode", "Places", "EYR OEIF/CIF: Most recent full: Overall effectiveness",
                            "EYR OEIF/CIF: Most recent full: Inspection date"]
                    miss = [c for c in need if c not in ix]
                    if miss:
                        raise LoadError(f"Childcare thiếu cột {miss}")
                continue

            def g(c):
                i = ix.get(c)
                v = r[i] if i is not None and i < len(r) else None
                return None if v in (None, "", "None", "REDACTED") else str(v).strip()
            rows.append([g(c) for c in CHILDCARE_FIELDS])
        if header is None:
            raise LoadError("Không thấy dòng tiêu đề sheet D1_Childcare_providers")
        if not 50_000 <= len(rows) <= 80_000:
            raise LoadError(f"Số cơ sở trông trẻ bất thường: {len(rows):,} (Ofsted ~60 nghìn)")
        cols = ["urn", "ptype", "subtype", "status", "name", "postcode", "places", "oeif", "oeif_date",
                "reif_date", "reif_curriculum", "reif_welfare"]
        st = ctx.stage_table("childcare", cols)
        ctx.copy_rows(st, cols, rows, "COPY childcare")
        public = ctx.scalar(f"SELECT count(*) FROM {st} WHERE name IS NOT NULL AND postcode IS NOT NULL")
        log.info("    %s/%s cơ sở có tên + postcode công khai (người giữ trẻ tại nhà bị ẩn)", f"{public:,}", f"{len(rows):,}")
        ctx.conn.execute("""INSERT INTO core.poi_type (code, parent_code, name_i18n, is_searchable, has_pages, url_segment)
            VALUES ('childcare', NULL, '{"en":"Childcare"}', false, false, NULL),
                   ('childcare.nursery', 'childcare', '{"en":"Nursery or pre-school"}', true, false, NULL)
            ON CONFLICT (code) DO NOTHING""")
        n = ctx.sql("upsert POI nhà trẻ", f"""
            INSERT INTO core.poi (market_code, type_code, external_scheme, external_id, name, geom, postal_code, area_id,
                                  status, attrs, dataset_id, release_id)
            SELECT 'uk', 'childcare.nursery', 'ofsted_ey', s.urn, initcap(s.name), p.geom, p.postcode, p.area_id, 'open',
                   jsonb_strip_nulls(jsonb_build_object('provider_type', s.ptype, 'subtype', s.subtype,
                       'places', nullif(s.places, '')::float::int)), 'uk.ofsted.childcare', %s
            FROM {st} s JOIN mkt_uk.postcode p ON p.postcode_compact = replace(upper(s.postcode), ' ', '')
            WHERE s.name IS NOT NULL AND p.geom IS NOT NULL AND s.status = 'Active'
              AND s.ptype IN ('Childcare on non-domestic premises', 'Childcare on domestic premises')
            ON CONFLICT (market_code, external_scheme, external_id) DO UPDATE SET name = EXCLUDED.name, geom = EXCLUDED.geom,
                postal_code = EXCLUDED.postal_code, area_id = EXCLUDED.area_id, status = 'open', attrs = EXCLUDED.attrs,
                release_id = EXCLUDED.release_id""", (ctx.release_id,))
        closed = ctx.sql("POI không còn trong danh sách → closed", """UPDATE core.poi SET status = 'closed'
            WHERE market_code='uk' AND external_scheme='ofsted_ey' AND release_id IS DISTINCT FROM %s AND status='open'""",
                         (ctx.release_id,))
        if not 15_000 <= n <= 40_000:
            raise LoadError(f"Số nhà trẻ nạp bất thường: {n:,}")
        r = ctx.sql("rating OEIF (khung cũ)", f"""
            INSERT INTO core.poi_rating (poi_id, scheme, aspect, grade_code, grade_label, grade_rank, grade_scale_max,
                                         rated_on, release_id)
            SELECT p.id, 'ofsted_ey', 'overall', g.code, g.label, g.rank, 4, s.oeif_date::timestamp::date, %s
            FROM {st} s JOIN core.poi p ON p.market_code='uk' AND p.external_scheme='ofsted_ey' AND p.external_id = s.urn
            -- MI của Ofsted ghi điểm dạng số (1–4, có thể '2.0') hoặc chữ
            CROSS JOIN LATERAL (SELECT r AS rank, (ARRAY['outstanding','good','requires_improvement','inadequate'])[r] AS code,
                                       (ARRAY['Outstanding','Good','Requires improvement','Inadequate'])[r] AS label
                                FROM (SELECT CASE WHEN s.oeif ~ '^[1-4]([.]0)?$' THEN s.oeif::numeric::int
                                                  WHEN lower(s.oeif) = 'outstanding' THEN 1 WHEN lower(s.oeif) = 'good' THEN 2
                                                  WHEN lower(s.oeif) = 'requires improvement' THEN 3
                                                  WHEN lower(s.oeif) = 'inadequate' THEN 4 END AS r) x) g
            WHERE g.rank IS NOT NULL AND s.oeif_date IS NOT NULL
            ON CONFLICT (poi_id, scheme, aspect, rated_on) DO UPDATE SET grade_code = EXCLUDED.grade_code,
                grade_rank = EXCLUDED.grade_rank, release_id = EXCLUDED.release_id""", (ctx.release_id,))
        r2 = ctx.sql("report card (khung mới, từ 11/2025)", f"""
            INSERT INTO core.poi_rating (poi_id, scheme, aspect, grade_code, grade_label, rated_on, release_id)
            SELECT p.id, 'ofsted_ey_report_card', a.aspect, a.grade, a.grade, s.reif_date::timestamp::date, %s
            FROM {st} s JOIN core.poi p ON p.market_code='uk' AND p.external_scheme='ofsted_ey' AND p.external_id = s.urn
            CROSS JOIN LATERAL (VALUES ('curriculum_teaching', s.reif_curriculum), ('welfare_wellbeing', s.reif_welfare)) a(aspect, grade)
            WHERE s.reif_date IS NOT NULL AND a.grade IS NOT NULL AND a.grade <> 'Not applicable'
            ON CONFLICT (poi_id, scheme, aspect, rated_on) DO UPDATE SET grade_code = EXCLUDED.grade_code,
                release_id = EXCLUDED.release_id""", (ctx.release_id,))
        vals = ctx.conn.execute(f"SELECT oeif, count(*) FROM {st} WHERE oeif IS NOT NULL GROUP BY 1 ORDER BY 2 DESC LIMIT 8").fetchall()
        log.info("    giá trị OEIF: %s", vals)
        good = ctx.scalar("""SELECT 100.0 * count(*) FILTER (WHERE grade_rank <= 2) / nullif(count(*), 0) FROM core.poi_rating
                             WHERE scheme = 'ofsted_ey' AND release_id = %s""", (ctx.release_id,))
        log.info("    %.1f%% nhà trẻ Good/Outstanding (Ofsted công bố ~96–97%%)", good or 0)
        if good is None or not 85 <= good <= 100:
            raise LoadError(f"Tỉ lệ Good/Outstanding bất thường: {good}")
        ctx.drop_stage("childcare")
        return {"providers": len(rows), "nurseries": n, "closed": closed, "ratings_oeif": r, "ratings_report_card": r2}


# ----------------------------------------------------- tuyển sinh cấp trường -----
OFFER_COLS = {  # cột stage → cột file cấp trường
    "places": "total_number_places_offered", "first_apps": "times_put_as_1st_preference",
    "second_apps": "times_put_as_2nd_preference", "third_apps": "times_put_as_3rd_preference",
    "any_apps": "times_put_as_any_preferred_school", "other_la_apps": "all_applications_from_another_LA",
    "first_offers": "number_1st_preference_offers", "second_offers": "number_2nd_preference_offers",
    "third_offers": "number_3rd_preference_offers", "preferred_offers": "number_preferred_offers",
    "other_la_offers": "offers_to_applicants_from_another_LA",
}


class SchoolOffersLoader(Loader):
    source = "dfe_school_offers"
    depends_on = ("gias",)
    description = "DfE applications & offers cấp trường → core.poi_metric_value uk.school.* (chỉ tiêu, đơn và offer theo nguyện vọng 1/2/3, LA khác)"

    def load(self, ctx: LoadContext) -> dict:
        z = ctx.file(r"^ees_.*\.zip$")
        members = zip_members(z, r"SchoolLevel.*\.csv$")
        if len(members) != 1:
            raise LoadError(f"Cần 1 file SchoolLevel*.csv, thấy {members}")
        rows = []
        with zipfile.ZipFile(z) as zf, zf.open(members[0]) as f:
            for r in csv.DictReader(io.TextIOWrapper(f, encoding="utf-8-sig", newline="")):
                tp = r.get("time_period", "")
                if not re.fullmatch(r"\d{6}", tp) or not r.get("school_urn"):
                    continue
                period = date(int(tp[:4]), 9, 1)
                rows.append((r["school_urn"].strip(), period, *(_num(r.get(c, "")) for c in OFFER_COLS.values()),
                             r.get("school_phase")))
        if len(rows) < 100_000:
            raise LoadError(f"Chỉ có {len(rows):,} dòng cấp trường (kỳ vọng nhiều năm × ~20 nghìn trường)")
        cols = ["urn", "period", *OFFER_COLS, "phase"]
        st = ctx.stage_table("offers", cols, {"period": "date", **{c: "float8" for c in OFFER_COLS}})
        ctx.copy_rows(st, cols, rows, "COPY offers")
        defs = {"uk.school.places_offered": ("Places offered at entry", "count", "s.places"),
                "uk.school.first_pref_apps": ("First-preference applications", "count", "s.first_apps"),
                "uk.school.first_pref_ratio": ("First-preference applications per place offered", "number",
                                               "round((s.first_apps / nullif(s.places, 0))::numeric, 2)"),
                "uk.school.first_pref_offer_rate": ("Share of first-preference applicants offered a place", "percent",
                                                    "round((100 * s.first_offers / nullif(s.first_apps, 0))::numeric, 1)"),
                "uk.school.second_pref_apps": ("Second-preference applications", "count", "s.second_apps"),
                "uk.school.third_pref_apps": ("Third-preference applications", "count", "s.third_apps"),
                "uk.school.any_pref_apps": ("Applications naming the school at any preference", "count", "s.any_apps"),
                "uk.school.other_la_apps": ("Applications from another local authority", "count", "s.other_la_apps"),
                "uk.school.first_pref_offers": ("Offers made to first-preference applicants", "count", "s.first_offers"),
                "uk.school.second_pref_offers": ("Offers made to second-preference applicants", "count", "s.second_offers"),
                "uk.school.third_pref_offers": ("Offers made to third-preference applicants", "count", "s.third_offers"),
                "uk.school.preferred_offers": ("Offers made to applicants who named the school", "count", "s.preferred_offers"),
                "uk.school.other_la_offers": ("Offers made to applicants from another local authority", "count", "s.other_la_offers")}
        for mid, (name, vt, _) in defs.items():
            register_metric(ctx, mid, category="education", dataset="uk.dfe.school_offers", value_type=vt,
                            decimals=2 if vt == "number" else (1 if vt == "percent" else 0),
                            unit="%" if vt == "percent" else None, direction=0, period_type="year", native_level=None,
                            aggregation="none", scope="uk.england", name=name, scorable=False,
                            description="DfE primary and secondary school applications and offers (national offer day), school level")
        _set_subject_poi(ctx, list(defs))
        ctx.sql("xoá giá trị cũ", "DELETE FROM core.poi_metric_value WHERE metric_id LIKE 'uk.school.%%' AND metric_id = ANY(%s)", (list(defs),))
        n = 0
        for mid, (_, _, expr) in defs.items():
            n += ctx.sql(mid, f"""
                INSERT INTO core.poi_metric_value (poi_id, metric_id, period, value_num, release_id)
                SELECT DISTINCT ON (p.id, s.period) p.id, %s, s.period, {expr}, %s FROM {st} s
                JOIN core.poi p ON p.market_code='uk' AND p.external_scheme='urn' AND p.external_id = s.urn
                WHERE {expr} IS NOT NULL ORDER BY p.id, s.period, s.places DESC NULLS LAST""", (mid, ctx.release_id))
        match = ctx.scalar(f"""SELECT count(DISTINCT s.urn)::float / count(DISTINCT s2.urn) FROM {st} s2
            LEFT JOIN core.poi p ON p.market_code='uk' AND p.external_scheme='urn' AND p.external_id = s2.urn
            LEFT JOIN {st} s ON s.urn = s2.urn AND p.id IS NOT NULL""")
        log.info("    %.1f%% URN khớp trường trong GIAS", match * 100)
        if match < 0.9:
            raise LoadError(f"Chỉ {match:.1%} URN khớp GIAS")
        bad = ctx.scalar(f"""SELECT count(*) FILTER (WHERE first_offers > first_apps OR second_offers > second_apps
                OR third_offers > third_apps OR first_apps > any_apps OR preferred_offers > places * 1.1 + 2)::float
                / count(*) FROM {st}""")
        log.info("    %.2f%% dòng có offer > đơn hoặc offer > chỉ tiêu", bad * 100)
        if bad > 0.01:
            raise LoadError(f"{bad:.1%} dòng có số offer/đơn mâu thuẫn")
        ctx.drop_stage("offers")
        return {"rows": len(rows), "values": n, "urn_match": round(match, 3)}


# -------------------------------------------------------------------- TfL -----
class TflStationsLoader(Loader):
    source = "tfl_stations"
    depends_on = ("naptan",)
    description = "TfL StopPoint → core.poi.attrs.tfl_zone / tfl_lines / tfl_modes cho ga NaPTAN (atco)"

    def load(self, ctx: LoadContext) -> dict:
        rows = {}
        for f in ctx.files(r"^stoppoints_.*\.json$"):
            mode = re.match(r"stoppoints_(.*)\.json$", f.name).group(1)
            for s in json.loads(f.read_text()).get("stopPoints", []):
                sid = s.get("stationNaptan") or s.get("naptanId")
                if not sid:
                    continue
                zone = next((a["value"] for a in s.get("additionalProperties", []) if a.get("key") == "Zone"), None)
                lines = sorted({ln.get("name") for ln in s.get("lines", []) if ln.get("name")})
                cur = rows.setdefault(sid, {"zone": None, "lines": set(), "modes": set(), "name": s.get("commonName")})
                cur["zone"] = cur["zone"] or zone
                cur["lines"].update(lines)
                cur["modes"].add(mode)
        if len(rows) < 400:
            raise LoadError(f"Chỉ có {len(rows)} ga TfL")
        st = ctx.stage_table("tfl", ["atco", "zone", "lines", "modes"], {"lines": "jsonb", "modes": "jsonb"})
        ctx.copy_rows(st, ["atco", "zone", "lines", "modes"],
                      [(k, v["zone"], json.dumps(sorted(v["lines"])), json.dumps(sorted(v["modes"]))) for k, v in rows.items()],
                      "COPY TfL")
        n = ctx.sql("gắn zone/tuyến vào POI ga", f"""
            UPDATE core.poi p SET attrs = p.attrs || jsonb_strip_nulls(jsonb_build_object('tfl_zone', s.zone,
                'tfl_lines', s.lines, 'tfl_modes', s.modes))
            FROM {st} s WHERE p.market_code='uk' AND p.external_scheme='atco'
              -- TfL dùng mã StopArea 940G…, POI NaPTAN của ga metro dùng mã 9400… cùng phần đuôi
              AND p.external_id IN (s.atco, '9400' || substr(s.atco, 5))""")
        with_zone = ctx.scalar(f"SELECT count(*) FROM {st} WHERE zone IS NOT NULL")
        log.info("    %s/%s ga TfL khớp POI NaPTAN, %s ga có zone", n, len(rows), with_zone)
        if n < 0.5 * len(rows) or with_zone < 300:
            raise LoadError(f"Ghép ga TfL ↔ NaPTAN kém: {n}/{len(rows)}, có zone {with_zone}")
        ctx.drop_stage("tfl")
        return {"stations": len(rows), "matched": n, "with_zone": with_zone}


GAP_TABLE_LOADERS = [RofrsLoader(), AffordabilityLoader(), PtalLoader(), FlyTippingLoader(), LhaLoader(), TribunalLoader(),
                     MpsLoader(), ChildcareLoader(), SchoolOffersLoader(), TflStationsLoader()]


# -------------------------------------------------------------------- CQC -----
CQC_FIELDS = ["Location ID", "Dormant (Y/N)", "Care home?", "Location Name", "Location ODS Code", "Care homes beds",
              "Location Primary Inspection Category", "Location Latest Overall Rating", "Publication Date",
              "Location Postal Code", "Location Latitude", "Location Longitude", "Provider Name", "Location Web Address"]
CQC_RANK = {"Outstanding": 1, "Good": 2, "Requires improvement": 3, "Inadequate": 4}


class CqcLoader(Loader):
    source = "cqc"
    depends_on = ("nhs_ods", "onspd")
    description = "CQC HSCA Active Locations → poi_rating 'cqc' cho POI NHS (mã ODS) + POI health.care_home (nhà dưỡng lão)"

    def load(self, ctx: LoadContext) -> dict:
        from pathlib import Path

        from common.ods import iter_ods_rows
        path = ctx.file(r"HSCA_Active_Locations\.ods$")
        header, rows = None, []
        for r in iter_ods_rows(Path(path), "HSCA_Active_Locations"):
            if header is None:
                header = r
                ix = {h: i for i, h in enumerate(header)}
                miss = [c for c in CQC_FIELDS if c not in ix]
                if miss:
                    raise LoadError(f"CQC thiếu cột {miss}")
                continue

            def g(c):
                i = ix[c]
                v = r[i] if i < len(r) else None
                return None if v in (None, "", "None", "*", "-") else str(v).strip()
            rows.append([g(c) for c in CQC_FIELDS])
        if not 40_000 <= len(rows) <= 70_000:
            raise LoadError(f"Số cơ sở CQC bất thường: {len(rows):,} (~55 nghìn)")
        cols = ["loc_id", "dormant", "care_home", "name", "ods", "beds", "category", "rating", "published", "postcode",
                "lat", "lon", "provider", "website"]
        st = ctx.stage_table("cqc", cols)
        ctx.copy_rows(st, cols, rows, "COPY CQC locations")
        ctx.check_values_in("xếp hạng", st, "rating", list(CQC_RANK) + [
            "Insufficient evidence to rate", "No rating", "Not applicable", "Not Rated", "No Approved Rating",
            "Evidence requirements partially / not yet scored"])
        ctx.check_range("vĩ độ", st, "nullif(lat, '')::float8", 49.8, 56)
        rank_sql = "CASE s.rating " + " ".join(f"WHEN '{k}' THEN {v}" for k, v in CQC_RANK.items()) + " END"
        r1 = ctx.sql("rating CQC cho POI NHS (mã ODS)", f"""
            INSERT INTO core.poi_rating (poi_id, scheme, aspect, grade_code, grade_label, grade_rank, grade_scale_max,
                                         rated_on, report_url, release_id)
            SELECT DISTINCT ON (p.id) p.id, 'cqc', 'overall', lower(replace(s.rating, ' ', '_')), s.rating, {rank_sql}, 4,
                   s.published::timestamp::date, 'https://www.cqc.org.uk/location/' || s.loc_id, %s
            FROM {st} s JOIN core.poi p ON p.market_code='uk' AND p.external_scheme='ods' AND p.external_id = s.ods
            WHERE s.rating IN ('Outstanding', 'Good', 'Requires improvement', 'Inadequate') AND s.published IS NOT NULL
            ORDER BY p.id, s.published DESC
            ON CONFLICT (poi_id, scheme, aspect, rated_on) DO UPDATE SET grade_code = EXCLUDED.grade_code,
                grade_label = EXCLUDED.grade_label, grade_rank = EXCLUDED.grade_rank, release_id = EXCLUDED.release_id""",
                     (ctx.release_id,))
        ctx.sql("is_current = bản mới nhất mỗi POI", """
            UPDATE core.poi_rating r SET is_current = (r.rated_on = m.max_on)
            FROM (SELECT poi_id, max(rated_on) max_on FROM core.poi_rating WHERE scheme = 'cqc' GROUP BY 1) m
            WHERE r.scheme = 'cqc' AND r.poi_id = m.poi_id""")
        gp = ctx.conn.execute("""SELECT count(*), 100.0 * count(*) FILTER (WHERE r.grade_rank <= 2) / nullif(count(*), 0)
            FROM core.poi_rating r JOIN core.poi p ON p.id = r.poi_id
            WHERE r.scheme = 'cqc' AND r.is_current AND p.type_code = 'health.gp'""").fetchone()
        log.info("    GP có xếp hạng CQC: %s, %.1f%% Good/Outstanding", f"{gp[0]:,}", gp[1] or 0)
        if gp[0] < 5000 or not 85 <= (gp[1] or 0) <= 99:  # ~6.200 practice England, ~95% Good/Outstanding
            raise LoadError(f"Xếp hạng CQC của GP bất thường: {gp}")

        ctx.conn.execute("""INSERT INTO core.poi_type (code, parent_code, name_i18n, is_searchable, has_pages, url_segment)
            VALUES ('health.care_home', 'health', '{"en":"Care home"}', true, false, NULL) ON CONFLICT (code) DO NOTHING""")
        n = ctx.sql("upsert POI nhà dưỡng lão", f"""
            INSERT INTO core.poi (market_code, type_code, external_scheme, external_id, name, geom, postal_code, status,
                                  attrs, dataset_id, release_id)
            SELECT 'uk', 'health.care_home', 'cqc', s.loc_id, s.name,
                   ST_SetSRID(ST_MakePoint(s.lon::float8, s.lat::float8), 4326), s.postcode, 'open',
                   jsonb_strip_nulls(jsonb_build_object('beds', nullif(s.beds, '')::float::int, 'provider', s.provider,
                       'category', s.category, 'website', s.website)), 'uk.cqc.directory', %s
            FROM {st} s WHERE s.care_home = 'Y' AND coalesce(s.dormant, 'N') = 'N' AND s.lat ~ '^-?[0-9.]+$' AND s.lon ~ '^-?[0-9.]+$'
            ON CONFLICT (market_code, external_scheme, external_id) DO UPDATE SET name = EXCLUDED.name, geom = EXCLUDED.geom,
                postal_code = EXCLUDED.postal_code, status = 'open', attrs = EXCLUDED.attrs, release_id = EXCLUDED.release_id""",
                    (ctx.release_id,))
        if not 10_000 <= n <= 20_000:  # ~14.500 nhà dưỡng lão England
            raise LoadError(f"Số nhà dưỡng lão bất thường: {n:,}")
        closed = ctx.sql("nhà dưỡng lão không còn → closed", """UPDATE core.poi SET status = 'closed'
            WHERE market_code='uk' AND external_scheme='cqc' AND release_id IS DISTINCT FROM %s AND status='open'""",
                         (ctx.release_id,))
        r2 = ctx.sql("rating CQC cho nhà dưỡng lão", f"""
            INSERT INTO core.poi_rating (poi_id, scheme, aspect, grade_code, grade_label, grade_rank, grade_scale_max,
                                         rated_on, report_url, release_id)
            SELECT p.id, 'cqc', 'overall', lower(replace(s.rating, ' ', '_')), s.rating, {rank_sql}, 4,
                   s.published::timestamp::date, 'https://www.cqc.org.uk/location/' || s.loc_id, %s
            FROM {st} s JOIN core.poi p ON p.market_code='uk' AND p.external_scheme='cqc' AND p.external_id = s.loc_id
            WHERE s.rating IN ('Outstanding', 'Good', 'Requires improvement', 'Inadequate') AND s.published IS NOT NULL
            ON CONFLICT (poi_id, scheme, aspect, rated_on) DO UPDATE SET grade_code = EXCLUDED.grade_code,
                grade_rank = EXCLUDED.grade_rank, release_id = EXCLUDED.release_id""", (ctx.release_id,))
        from .poi import attach_nearest_postcode
        attach_nearest_postcode(ctx, "cqc")
        ctx.drop_stage("cqc")
        return {"locations": len(rows), "nhs_ratings": r1, "gp_rated": gp[0], "care_homes": n, "closed": closed,
                "care_home_ratings": r2}


GAP_TABLE_LOADERS.append(CqcLoader())


# ------------------------------------------------- xả tràn, sự cố ô nhiễm -----
def ngr_to_en(ngr: str | None) -> tuple[float, float] | None:
    """Toạ độ lưới OS dạng chữ ('SP6419046470', 'NY 3634 6325') → (easting, northing) tâm ô (m)."""
    if not ngr:
        return None
    s = re.sub(r"\s+", "", str(ngr)).upper()
    m = re.fullmatch(r"([A-HJ-Z]{2})(\d+)", s)
    if not m or len(m.group(2)) % 2 or not 2 <= len(m.group(2)) <= 10:
        return None
    a, b = (ord(c) - 65 - (c > "I") for c in m.group(1))  # bỏ chữ I
    e100 = ((a - 2) % 5) * 5 + (b % 5)
    n100 = (19 - (a // 5) * 5) - (b // 5)
    digits = m.group(2)
    k = len(digits) // 2
    unit = 10 ** (5 - k)
    e = e100 * 100_000 + int(digits[:k]) * unit + unit / 2
    n = n100 * 100_000 + int(digits[k:]) * unit + unit / 2
    if not (0 <= e < 700_000 and 0 <= n < 1_300_000):
        return None
    return e, n


def _poi_types(ctx: LoadContext, rows: list[tuple[str, str | None, str]]) -> None:
    for code, parent, name in rows:
        ctx.conn.execute("""INSERT INTO core.poi_type (code, parent_code, name_i18n, is_searchable, has_pages, url_segment)
            VALUES (%s, %s, jsonb_build_object('en', %s::text), false, false, NULL) ON CONFLICT (code) DO NOTHING""",
                         (code, parent, name))


def _hours(v) -> float | None:
    if v is None:
        return None
    s = str(v).strip()
    m = re.fullmatch(r"(\d+):(\d{2}):(\d{2})", s)
    if m:
        return int(m.group(1)) + int(m.group(2)) / 60 + int(m.group(3)) / 3600
    return _num(s)


class StormOverflowLoader(Loader):
    source = "ea_storm_overflows"
    description = "EA EDM: điểm xả tràn nước thải (năm mới nhất) → core.poi env.storm_overflow (attrs: số lần, số giờ xả)"

    def load(self, ctx: LoadContext) -> dict:
        import openpyxl
        zips = sorted(ctx.files(r"^EDM_20\d\d_Storm_Overflow_Annual_Return\.zip$"))
        latest = zips[-1]
        year = int(re.search(r"EDM_(\d{4})_", latest.name).group(1))
        member = zip_members(latest, r"all water and sewerage companies\.xlsx$")
        if len(member) != 1:
            raise LoadError(f"{latest.name}: cần 1 file tổng hợp mọi công ty, thấy {member}")
        with zipfile.ZipFile(latest) as zf:
            data = io.BytesIO(zf.read(member[0]))
        ws = openpyxl.load_workbook(data, read_only=True, data_only=True).worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = None
        for r in it:
            if r and r[0] == "Unique ID":
                header = [str(h or "").split("\n")[0].strip() for h in r]
                break
        if not header:
            raise LoadError("Không thấy dòng tiêu đề 'Unique ID'")
        ix = {h: i for i, h in enumerate(header)}
        need = ["Unique ID", "Water Company Name", "Site Name", "Outlet Discharge NGR", "Counted spills using 12-24h count method",
                "Total Duration (hh:mm:ss) all spills prior to processing through 12-24h count method",
                "Receiving Water / Environment (common name)", "Long-term average spill count", "Bathing Water(s) (only populate for storm overflows with a",
                "Storm Discharge Asset Type"]
        def col(prefix):
            hits = [i for h, i in ix.items() if h.startswith(prefix[:30])]
            return hits[0] if hits else None
        cidx = {k: col(k) for k in need}
        miss = [k for k, v in cidx.items() if v is None]
        if miss:
            raise LoadError(f"EDM thiếu cột {miss}; có {header}")
        rows, bad_ngr = [], 0
        for r in it:
            uid = r[cidx["Unique ID"]] if r else None
            if not uid:
                continue
            en = ngr_to_en(r[cidx["Outlet Discharge NGR"]])
            if not en:
                bad_ngr += 1
                continue
            rows.append((str(uid).strip(), r[cidx["Site Name"]], r[cidx["Water Company Name"]], en[0], en[1],
                         _num(r[cidx["Counted spills using 12-24h count method"]]),
                         _hours(r[cidx["Total Duration (hh:mm:ss) all spills prior to processing through 12-24h count method"]]),
                         _num(r[cidx["Long-term average spill count"]]), r[cidx["Receiving Water / Environment (common name)"]],
                         r[cidx["Bathing Water(s) (only populate for storm overflows with a"]], r[cidx["Storm Discharge Asset Type"]]))
        log.info("    %s điểm xả năm %s, %s NGR lỗi", f"{len(rows):,}", year, bad_ngr)
        if not 12_000 <= len(rows) <= 16_000 or bad_ngr > 0.02 * len(rows):
            raise LoadError(f"Số điểm xả bất thường: {len(rows)} (NGR lỗi {bad_ngr})")
        cols = ["uid", "site", "company", "e", "n", "spills", "hours", "lt_avg", "water", "bathing", "asset"]
        st = ctx.stage_table("edm", cols, {"e": "float8", "n": "float8", "spills": "float8", "hours": "float8", "lt_avg": "float8"})
        ctx.copy_rows(st, cols, [tuple(None if v in ("", "None") else v for v in r) for r in rows], "COPY EDM")
        ctx.check_range("số lần xả", st, "spills", 0, 400)
        tot = ctx.conn.execute(f"SELECT sum(spills), sum(hours) FROM {st}").fetchone()
        log.info("    tổng %s lần xả, %s giờ (năm %s)", f"{tot[0]:,.0f}", f"{tot[1]:,.0f}", year)
        if not 200_000 <= tot[0] <= 700_000:  # England ~ 300–480 nghìn lần/năm (2020–2025)
            raise LoadError(f"Tổng số lần xả bất thường: {tot[0]}")
        _poi_types(ctx, [("environment", None, "Environment"), ("environment.storm_overflow", "environment", "Storm overflow")])
        n = ctx.sql("upsert POI điểm xả", f"""
            INSERT INTO core.poi (market_code, type_code, external_scheme, external_id, name, geom, status, attrs,
                                  dataset_id, release_id)
            SELECT DISTINCT ON (uid) 'uk', 'environment.storm_overflow', 'edm', uid, coalesce(initcap(site), uid),
                   ST_Transform(ST_SetSRID(ST_MakePoint(e, n), 27700), 4326), 'open',
                   jsonb_strip_nulls(jsonb_build_object('year', %s, 'spills', spills, 'spill_hours', round(hours::numeric, 1),
                       'long_term_avg_spills', round(lt_avg::numeric, 1), 'company', company, 'receiving_water', water,
                       'bathing_water', nullif(bathing, 'Not Applicable'), 'asset_type', asset)), 'uk.ea.storm_overflows', %s
            FROM {st} ORDER BY uid
            ON CONFLICT (market_code, external_scheme, external_id) DO UPDATE SET name = EXCLUDED.name, geom = EXCLUDED.geom,
                status = 'open', attrs = EXCLUDED.attrs, release_id = EXCLUDED.release_id""", (year, ctx.release_id))
        ctx.sql("điểm xả không còn → closed", """UPDATE core.poi SET status = 'closed'
            WHERE market_code='uk' AND external_scheme='edm' AND release_id IS DISTINCT FROM %s AND status='open'""", (ctx.release_id,))
        from .poi import attach_nearest_postcode
        attach_nearest_postcode(ctx, "edm")
        ctx.drop_stage("edm")
        return {"overflows": n, "year": year, "spills": tot[0], "hours": round(tot[1])}


class PollutionIncidentsLoader(Loader):
    source = "ea_pollution_incidents"
    description = "EA NIRS: sự cố ô nhiễm môi trường mức 1–2 (Access .accdb) → core.poi environment.pollution_incident"

    def load(self, ctx: LoadContext) -> dict:
        import tempfile
        from collections import defaultdict

        from access_parser import AccessParser
        z = ctx.file(r"\.zip$")
        acc = zip_members(z, r"\.accdb$")
        with tempfile.TemporaryDirectory() as d, zipfile.ZipFile(z) as zf:
            path = zf.extract(acc[0], d)
            db = AccessParser(path)
            inc = db.parse_table("EP_Incidents_Nirs2")
            pol = db.parse_table("EP_Pollutants_Nirs2")
        pollutants = defaultdict(list)
        for nid, ptype, p in zip(pol["NOT_ID"], pol["POLL_TYPE"], pol["POLLUTANT"]):
            pollutants[nid].append(p or ptype)
        rows = []
        for i in range(len(inc["NOT_ID"])):
            x, y = inc["X_CONF"][i], inc["Y_CONF"][i]
            if not x or not y:
                continue
            cats = {k: inc[k][i] for k in ("EIL_AIR", "EIL_LAND", "EIL_WATER")}
            worst = min((int(m.group(1)) for v in cats.values() if v and (m := re.search(r"Category (\d)", v))), default=None)
            rows.append((str(inc["NOT_ID"][i]), str(inc["NOT_DATE"][i])[:10], float(x), float(y), worst,
                         json.dumps({k.lower(): v for k, v in cats.items()}), json.dumps(sorted(set(pollutants[inc["NOT_ID"][i]]))[:10]),
                         inc["DISTRICT"][i]))
        if not 10_000 <= len(rows) <= 30_000:
            raise LoadError(f"Số sự cố bất thường: {len(rows)}")
        cols = ["nid", "dt", "x", "y", "worst", "impacts", "pollutants", "district"]
        st = ctx.stage_table("nirs", cols, {"x": "float8", "y": "float8", "worst": "int", "impacts": "jsonb", "pollutants": "jsonb"})
        ctx.copy_rows(st, cols, rows, "COPY sự cố")
        ctx.check_range("easting", st, "x", 50_000, 700_000)
        ctx.check_range("mức nặng nhất", st, "worst", 1, 4)
        span = ctx.conn.execute(f"SELECT min(dt), max(dt) FROM {st}").fetchone()
        log.info("    %s sự cố, %s → %s", f"{len(rows):,}", *span)
        _poi_types(ctx, [("environment", None, "Environment"),
                         ("environment.pollution_incident", "environment", "Serious pollution incident")])
        n = ctx.sql("upsert POI sự cố (mức 1–2)", f"""
            INSERT INTO core.poi (market_code, type_code, external_scheme, external_id, name, geom, status, opened_on, attrs,
                                  dataset_id, release_id)
            SELECT 'uk', 'environment.pollution_incident', 'ea_nirs', nid, 'Pollution incident ' || dt,
                   ST_Transform(ST_SetSRID(ST_MakePoint(x, y), 27700), 4326), 'closed', dt::date,
                   jsonb_build_object('worst_category', worst, 'impacts', impacts, 'pollutants', pollutants), 'uk.ea.pollution_incidents', %s
            FROM {st} WHERE worst <= 2
            ON CONFLICT (market_code, external_scheme, external_id) DO UPDATE SET geom = EXCLUDED.geom, attrs = EXCLUDED.attrs,
                release_id = EXCLUDED.release_id""", (ctx.release_id,))
        if n < 5000:
            raise LoadError(f"Chỉ {n} sự cố mức 1–2")
        ctx.drop_stage("nirs")
        return {"incidents": len(rows), "serious": n, "from": span[0], "to": span[1]}


GAP_TABLE_LOADERS += [StormOverflowLoader(), PollutionIncidentsLoader()]


# ------------------------------------------------- dân số giữa năm (LSOA) -----
class LsoaPopulationLoader(Loader):
    source = "ons_lsoa_population"
    depends_on = ("boundaries_lsoa21", "census2021")
    description = "ONS mid-year estimates LSOA 2021 (mọi năm trong file) → uk.pop.mid_year (tổng + nhóm tuổi trong value_json)"

    def load(self, ctx: LoadContext) -> dict:
        import openpyxl
        wb = openpyxl.load_workbook(ctx.file(r"\.xlsx$"), read_only=True, data_only=True)
        rows, years = [], []
        for ws in wb.worksheets:
            m = re.search(r"Mid-(\d{4})", ws.title)
            if not m:
                continue
            year = int(m.group(1))
            it = ws.iter_rows(values_only=True)
            header = None
            for r in it:
                cells = [str(c).strip() if c is not None else "" for c in r]
                if "LSOA 2021 Code" in cells and "Total" in cells:
                    header = cells
                    break
            if not header:
                continue
            i_code, i_tot = header.index("LSOA 2021 Code"), header.index("Total")
            ages = {i: int(re.search(r"(\d+)$", h).group(1)) for i, h in enumerate(header) if re.fullmatch(r"[FM]\d+", h)}
            n0 = len(rows)
            for r in it:
                if not r or not r[i_code] or not re.match(r"^[EW]010\d{5}$", str(r[i_code])):
                    continue
                tot = _num(r[i_tot])
                bands = {"0-15": 0.0, "16-64": 0.0, "65+": 0.0}
                for i, a in ages.items():
                    v = _num(r[i]) or 0
                    bands["0-15" if a <= 15 else "16-64" if a <= 64 else "65+"] += v
                rows.append(("lsoa", r[i_code], "uk.pop.mid_year", date(year, 6, 30), tot, json.dumps(bands) if ages else None))
            if len(rows) - n0:
                years.append(year)
                log.info("    sheet %s: %s LSOA", ws.title, f"{len(rows) - n0:,}")
        if not years:
            raise LoadError("Không thấy sheet 'Mid-YYYY' có cột 'LSOA 2021 Code' và 'Total'")
        st = ctx.stage_table("pop", ["level", "code", "metric_id", "period", "value_num", "value_json"],
                             {"period": "date", "value_num": "float8", "value_json": "jsonb"})
        ctx.copy_rows(st, ["level", "code", "metric_id", "period", "value_num", "value_json"], rows, "COPY dân số")
        ctx.check_range("dân số LSOA", st, "value_num", 300, 15_000)
        latest = max(years)
        tot = ctx.scalar(f"SELECT sum(value_num) FROM {st} WHERE period = %s", (date(latest, 6, 30),))
        log.info("    E&W giữa năm %s: %s người", latest, f"{tot:,.0f}")
        if not 60_000_000 <= tot <= 64_000_000:  # E&W giữa 2024 ~ 61,8 triệu
            raise LoadError(f"Tổng dân số E&W bất thường: {tot:,.0f}")
        register_metric(ctx, "uk.pop.mid_year", category="people", dataset="uk.ons.lsoa_population", value_type="count",
                        decimals=0, direction=0, period_type="year", native_level="lsoa", aggregation="sum", scope="uk.ew",
                        name="Population (mid-year estimate)", scorable=False,
                        description="ONS small area population estimates; value_json = residents by age band")
        ctx.sql("xoá giá trị cũ", "DELETE FROM core.metric_value WHERE market_code='uk' AND metric_id='uk.pop.mid_year'")
        n = ctx.sql("INSERT dân số", f"""
            INSERT INTO core.metric_value (market_code, area_id, metric_id, period, value_num, value_json, release_id)
            SELECT 'uk', a.id, s.metric_id, s.period, s.value_num, s.value_json, %s FROM {st} s
            JOIN core.geo_area a ON a.market_code='uk' AND a.level_code='lsoa' AND a.code = s.code""", (ctx.release_id,))
        ctx.expect("giá trị khớp LSOA", n, len(rows), tolerance=0.001)
        ctx.drop_stage("pop")
        return {"values": n, "years": years, "latest_total": round(tot)}


GAP_TABLE_LOADERS.append(LsoaPopulationLoader())


# ------------------------------------------------------------ hồ chứa -----
class ReservoirLoader(Loader):
    source = "ea_reservoir_flood"
    depends_on = ("onspd", "census2021")
    description = "EA Reservoir Flood Extents (mặt nạ 10 m) tại postcode → uk.zone.reservoir_{wet,dry}_share theo LSOA (England)"

    def load(self, ctx: LoadContext) -> dict:
        import itertools

        import numpy as np
        import psycopg
        from PIL import Image

        from common.config import database_url
        T = 10_000
        tiles = {}
        for p in ctx.files(r"^reservoir_\d{3}_\d{3}\.png$"):
            m = re.match(r"reservoir_(\d{3})_(\d{3})\.png$", p.name)
            tiles[(int(m.group(1)), int(m.group(2)))] = p
        if len(tiles) < 100:
            raise LoadError(f"Chỉ có {len(tiles)} ô hồ chứa")
        st = ctx.stage_table("res_pc", ["postcode", "v"], {"v": "smallint"})
        rconn = psycopg.connect(database_url())
        cur = rconn.cursor(name="pcs")
        cur.itersize = 50_000
        cur.execute(f"""SELECT postcode, ST_X(g), ST_Y(g), floor(ST_X(g) / {T})::int, floor(ST_Y(g) / {T})::int
            FROM (SELECT postcode, ST_Transform(geom, 27700) g FROM mkt_uk.postcode
                  WHERE country_code = 'E92000001' AND status = 'live' AND geom IS NOT NULL) a ORDER BY 4, 5""")

        def rows():
            for (tx, ty), grp in itertools.groupby(cur, key=lambda r: (r[3], r[4])):
                pcs = list(grp)
                path = tiles.get((tx, ty))
                if path is None:
                    for p in pcs:
                        yield (p[0], 0)
                    continue
                a = np.asarray(Image.open(path))
                sx, _, _, nsy, x0, y0 = [float(x) for x in path.with_suffix(".pgw").read_text().split()]
                h, w = a.shape
                col = np.clip(np.rint((np.array([p[1] for p in pcs]) - x0) / sx).astype(int), 0, w - 1)
                row = np.clip(np.rint((np.array([p[2] for p in pcs]) - y0) / nsy).astype(int), 0, h - 1)
                for p, v in zip(pcs, a[row, col].tolist()):
                    yield (p[0], v)
        try:
            n = ctx.copy_rows(st, ["postcode", "v"], rows(), "lấy mẫu vùng ngập hồ chứa tại postcode")
        finally:
            cur.close()
            rconn.close()
        ctx.sql("dân theo postcode", f"""
            CREATE TEMP TABLE res_pop ON COMMIT DROP AS
            SELECT p.lsoa21_code, s.v, mv.value_num / count(*) OVER (PARTITION BY p.oa21_code) AS pop
            FROM {st} s JOIN mkt_uk.postcode p USING (postcode)
            JOIN core.geo_area oa ON oa.market_code='uk' AND oa.level_code='oa' AND oa.code = p.oa21_code
            JOIN core.metric_value mv ON mv.market_code='uk' AND mv.area_id = oa.id AND mv.metric_id='uk.census21.ts001'""")
        eng = ctx.conn.execute("""SELECT 100.0 * sum(pop) FILTER (WHERE v & 2 > 0) / sum(pop),
                                         100.0 * sum(pop) FILTER (WHERE v & 1 > 0) / sum(pop) FROM res_pop""").fetchone()
        log.info("    %% dân England trong vùng ngập hồ chứa: ngày mưa %.2f%%, ngày khô %.2f%%", *eng)
        if not (0.5 <= eng[0] <= 20 and eng[1] <= eng[0] * 1.2):
            raise LoadError(f"Tỉ lệ dân vùng ngập hồ chứa bất thường: {eng}")
        period = date.fromisoformat(ctx.version)
        defs = {"reservoir_wet": ("Residents in the reservoir flood extent (wet day: reservoir failure plus river flooding)", 2, True),
                "reservoir_dry": ("Residents in the reservoir flood extent (dry day: reservoir failure only)", 1, False)}
        for k, (name, _, scorable) in defs.items():
            register_metric(ctx, f"uk.zone.{k}_share", category="risk", dataset="uk.ea.reservoir_flood", value_type="percent",
                            unit="%", decimals=1, direction=-1, period_type="snapshot", native_level="lsoa",
                            aggregation="weighted_mean", scope="uk.england", name=name, scorable=scorable,
                            description="EA Reservoir Flood Extents (worst-case, unlikely event) sampled at each postcode centre")
        ids = [f"uk.zone.{k}_share" for k in defs]
        ctx.sql("xoá giá trị cũ", "DELETE FROM core.metric_value WHERE market_code='uk' AND metric_id = ANY(%s)", (ids,))
        ins = ctx.sql("tỉ lệ dân theo LSOA", """
            INSERT INTO core.metric_value (market_code, area_id, metric_id, period, value_num, sample_size, release_id)
            SELECT 'uk', a.id, 'uk.zone.' || d.k || '_share', %s,
                   round((100.0 * coalesce(sum(r.pop) FILTER (WHERE r.v & d.bit > 0), 0) / sum(r.pop))::numeric, 2), count(*), %s
            FROM res_pop r JOIN core.geo_area a ON a.market_code='uk' AND a.level_code='lsoa' AND a.code = r.lsoa21_code
            CROSS JOIN (VALUES ('reservoir_wet', 2), ('reservoir_dry', 1)) d(k, bit)
            GROUP BY a.id, d.k HAVING sum(r.pop) > 0""", (period, ctx.release_id))
        ctx.drop_stage("res_pc")
        return {"postcodes": n, "values": ins, "england_pct": {"wet": round(eng[0], 2), "dry": round(eng[1], 2)}}


GAP_TABLE_LOADERS.append(ReservoirLoader())


# ---------------------------------------------------------- lưu lượng xe -----
class TrafficLoader(Loader):
    source = "dft_traffic"
    description = "DfT AADF: năm mới nhất của mỗi điểm đếm → core.poi transport.traffic_count (xe/ngày, xe tải, xe đạp)"
    COLS = ["count_point_id", "year", "road_name", "road_category", "road_type", "latitude", "longitude",
            "estimation_method", "pedal_cycles", "all_HGVs", "all_motor_vehicles"]

    def load(self, ctx: LoadContext) -> dict:
        z = ctx.file(r"\.zip$")
        member = zip_members(z, r"aadf\.csv$")
        if len(member) != 1:
            raise LoadError(f"Cần 1 file *aadf.csv, thấy {member}")
        latest = {}
        with zipfile.ZipFile(z) as zf, zf.open(member[0]) as f:
            rd = csv.DictReader(io.TextIOWrapper(f, encoding="utf-8-sig", newline=""))
            miss = [c for c in self.COLS if c not in rd.fieldnames]
            if miss:
                raise LoadError(f"AADF thiếu cột {miss}")
            n_rows = 0
            for r in rd:
                n_rows += 1
                cp, y = r["count_point_id"], int(r["year"])
                if cp not in latest or y > latest[cp][1]:
                    latest[cp] = (cp, y, r["road_name"], r["road_category"], r["road_type"], _num(r["latitude"]),
                                  _num(r["longitude"]), r["estimation_method"], _num(r["pedal_cycles"]),
                                  _num(r["all_HGVs"]), _num(r["all_motor_vehicles"]))
        rows = [v for v in latest.values() if v[5] is not None and v[6] is not None]
        log.info("    %s dòng AADF, %s điểm đếm", f"{n_rows:,}", f"{len(rows):,}")
        if not 30_000 <= len(rows) <= 80_000:
            raise LoadError(f"Số điểm đếm bất thường: {len(rows)}")
        cols = ["cp", "year", "road", "cat", "rtype", "lat", "lon", "method", "cycles", "hgv", "motor"]
        st = ctx.stage_table("aadf", cols, {"year": "int", "lat": "float8", "lon": "float8", "cycles": "float8",
                                            "hgv": "float8", "motor": "float8"})
        ctx.copy_rows(st, cols, rows, "COPY điểm đếm")
        ctx.check_range("xe/ngày", st, "motor", 0, 300_000)
        ctx.check_range("vĩ độ", st, "lat", 49.8, 61)
        top = ctx.scalar(f"SELECT max(motor) FROM {st} WHERE road LIKE 'M25%%'")
        log.info("    M25 lớn nhất: %s xe/ngày; năm mới nhất %s", f"{top:,.0f}" if top else None, ctx.scalar(f"SELECT max(year) FROM {st}"))
        if not top or top < 100_000:  # M25 đoạn đông nhất ~200 nghìn xe/ngày
            raise LoadError(f"Lưu lượng M25 bất thường: {top}")
        _poi_types(ctx, [("transport.traffic_count", "transport", "Traffic count point")])
        n = ctx.sql("upsert POI điểm đếm", f"""
            INSERT INTO core.poi (market_code, type_code, external_scheme, external_id, name, geom, status, attrs,
                                  dataset_id, release_id)
            SELECT 'uk', 'transport.traffic_count', 'dft_cp', cp, coalesce(nullif(road, ''), 'Count point ' || cp),
                   ST_SetSRID(ST_MakePoint(lon, lat), 4326), 'open',
                   jsonb_strip_nulls(jsonb_build_object('year', year, 'aadf_motor_vehicles', motor, 'aadf_hgv', hgv,
                       'aadf_pedal_cycles', cycles, 'road_category', cat, 'road_type', rtype, 'estimation', method)),
                   'uk.dft.traffic_aadf', %s
            FROM {st}
            ON CONFLICT (market_code, external_scheme, external_id) DO UPDATE SET name = EXCLUDED.name, geom = EXCLUDED.geom,
                attrs = EXCLUDED.attrs, release_id = EXCLUDED.release_id""", (ctx.release_id,))
        from .poi import attach_nearest_postcode
        attach_nearest_postcode(ctx, "dft_cp")
        ctx.drop_stage("aadf")
        return {"rows": n_rows, "count_points": n}


GAP_TABLE_LOADERS.append(TrafficLoader())


# ------------------------------------------------- công trình chống lũ -----
class FloodDefenceLoader(Loader):
    source = "ea_flood_defences"
    description = "EA AIMS Spatial Flood Defences (đường) → mkt_uk.flood_defence (loại, tình trạng, mức bảo vệ)"
    DDL = """
CREATE TABLE IF NOT EXISTS mkt_uk.flood_defence (
    asset_id text PRIMARY KEY, name text, asset_type text, protection_type text, current_condition text,
    target_condition text, design_sop int, current_sop int, maintainer text,
    geom geometry(MultiLineString, 4326) NOT NULL, release_id bigint REFERENCES ops.dataset_release (id)
);
CREATE INDEX IF NOT EXISTS flood_defence_geom ON mkt_uk.flood_defence USING gist (geom);
COMMENT ON TABLE mkt_uk.flood_defence IS 'EA AIMS: đê, tường, công trình chống lũ (England). current_condition 1–5 (1 = rất tốt, 5 = rất kém); *_sop = mức bảo vệ (chu kỳ lặp năm).'"""

    def load(self, ctx: LoadContext) -> dict:
        z = ctx.file(r"\.geojson\.zip$")
        m = zip_members(z, r"\.geojson$")
        ctx.ogr_load(f"/vsizip/{z}/{m[0]}", "aims", extra=["-dim", "XY", "-nlt", "MULTILINESTRING"])
        n_src = ctx.scalar(f"SELECT count(*) FROM {STAGE_T}")
        if not 100_000 <= n_src <= 200_000:
            raise LoadError(f"Số công trình chống lũ bất thường: {n_src:,}")
        _ddl(ctx, self.DDL)
        ctx.sql("TRUNCATE flood_defence", "TRUNCATE mkt_uk.flood_defence")
        n = ctx.sql("INSERT flood_defence", f"""
            INSERT INTO mkt_uk.flood_defence (asset_id, name, asset_type, protection_type, current_condition,
                target_condition, design_sop, current_sop, maintainer, geom, release_id)
            SELECT DISTINCT ON (asset_id) asset_id, nullif(asset_name, ''), asset_sub_type, protection_type,
                   current_condition, target_condition, design_sop, current_sop, asset_maintainer, ST_Multi(geom), %s
            FROM {STAGE_T} WHERE geom IS NOT NULL AND asset_id IS NOT NULL ORDER BY asset_id""", (ctx.release_id,))
        km = ctx.scalar("SELECT sum(ST_Length(ST_Transform(geom, 27700))) / 1000 FROM mkt_uk.flood_defence")
        log.info("    %s công trình, %.0f km", f"{n:,}", km)
        if not 5_000 <= km <= 60_000:  # EA quản lý ~ 9.000 km đê sông + biển, AIMS gồm cả tường/kênh nhỏ
            raise LoadError(f"Tổng chiều dài công trình bất thường: {km:.0f} km")
        ctx.drop_stage("aims")
        return {"assets": n, "km": round(km)}


STAGE_T = "stage_uk.aims"
GAP_TABLE_LOADERS.append(FloodDefenceLoader())
